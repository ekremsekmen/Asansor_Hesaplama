# -*- coding: utf-8 -*-
"""
ŞABLON GÜNCELLEME  —  PALANGA VERİM DÜŞÜŞÜ ( Δη = 0,10 ) KALDIRILMASI

Motor tarafında Δη kaldırıldı ( bkz. engine/ortak/ofis.py ).  Avan çalışma
kitabı aynı kuralı KENDİ formüllerinde taşıyordu;  yamalanmazsa ekran ile
indirilen dosya ayrışır — ekran η = 0,85 ile, kitap η′ = 0,75 ile hesaplar.

SATIR / SÜTUN KAYDIRMAZ.  Silinen satırlar boşaltılır, formüller yeniden
yazılır;  başka hiçbir formülün adresi değişmez.

Yapılanlar
  1  1..4 NOLU ASANSÖR :  η satırı ( 10 ) "toplam sistem verimi" olur
  2  1..4 NOLU ASANSÖR :  η′ satırı ( 11 ) tamamen boşaltılır
  3  1..4 NOLU ASANSÖR :  N formülü ( C13/C14/E14 ) E11 yerine E10 okur
  4  SABİTLER          :  Δη satırı ( 11 ) boşaltılır
  5  GİRİŞ             :  "Girilen η TOPLAM sistem verimi mi ?" satırı ( 56 )
                          boşaltılır, açılır listesi kaldırılır, 58. satırdaki
                          açıklama metni güncellenir

Çalıştırma:  python3 araclar/sablon_palanga_verim.py
"""
import os
import shutil
import sys

import openpyxl

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AVAN = os.path.join(KOK, "templates", "ASANSOR_AVAN_HESAPLARI.xlsx")

#  Asansör sayfası  ->  GİRİŞ'te o asansörün kolonu
SAYFA_KOLON = {"1 NOLU ASANSÖR": "C", "2 NOLU ASANSÖR": "D",
               "3 NOLU ASANSÖR": "E", "4 NOLU ASANSÖR": "F"}

ETA_ETIKET = "Toplam sistem verimi  ( askı / palanga kaybı DÂHİL )"
GIRIS_ACIKLAMA = (
    "Bir projede dişlili ve dişlisiz makine ya da 1:1 ve 2:1 askı birlikte "
    "kullanılabildiği için bu iki değer — toplam sistem verimi η gibi — "
    "asansöre özeldir. Boş bırakılan kolon SABİTLER B'deki ofis standardını "
    "kullanır.   η TOPLAM SİSTEM VERİMİDİR: askı ( palanga ), kasnak ve "
    "makine kayıpları içindedir; imalatçı kataloğundaki değer ( EN 81-20/50 "
    "şablonlarında η_ins ) doğrudan girilir. MMO/697 §2.4'ün Δη = 0,10 "
    "palanga düşüşü KALDIRILMIŞTIR — makara kaybı çarpımsaldır ve sabit bir "
    "sayı çıkarmak dişli ile dişlisiz makineyi farklı oranda cezalandırıyordu."
)


def bosalt(ws, satir, sutunlar="ABCDEFG"):
    for c in sutunlar:
        ws[f"{c}{satir}"] = None


def yama(yol):
    wb = openpyxl.load_workbook(yol)

    #  1-3  Asansör paftaları
    for sayfa, kol in SAYFA_KOLON.items():
        ws = wb[sayfa]
        ws[f"C10"] = ETA_ETIKET
        bosalt(ws, 11)
        ws["C13"] = "N   =   ( 1 − q ) · Q · V   /   ( 102 · η )"
        for hucre in ("C14", "E14"):
            f = ws[hucre].value
            if not isinstance(f, str):
                continue
            #  E11 -> E10  ( yalnız BU sayfanın kendi hücre atfı;  GİRİŞ! ve
            #  SABİTLER! ön ekli adresler zaten "E11" içermez )
            ws[hucre] = f.replace("E11", "E10")

    #  4  SABİTLER — Δη satırı
    bosalt(wb["SABİTLER"], 11, "ABC")

    #  5  GİRİŞ — toplam verim anahtarı ve yardım metinleri
    g = wb["GİRİŞ"]
    bosalt(g, 56, "ABCDEFGH")          # G/H yardım sütunları da temizlenir
    kalan = [dv for dv in g.data_validations.dataValidation
             if str(dv.sqref) != "C56:F56"]
    g.data_validations.dataValidation = kalan
    g["B29"] = "Toplam sistem verimi"
    g["H29"] = ("Askı ( palanga ), kasnak ve makine kayıpları DÂHİL tek verim. "
                "Ofis kabulü: Dişlisiz 0,85 / Dişli 0,50. İmalatçı kataloğundaki "
                "toplam sistem verimini ( EN 81-20/50 şablonlarında η_ins ) "
                "doğrudan girin.")
    g["H53"] = ("Asansörün kendi özelliğidir. Motor GÜCÜ askı oranından "
                "BAĞIMSIZDIR ve askı oranı verime de girmez — Δη = 0,10 "
                "palanga düşüşü kaldırılmıştır. Boş = SABİTLER B")
    g["H55"] = ("Dişlisiz  η = 0,85   /   Dişli  η = 0,50   ( ofis kabulü, "
                "TOPLAM sistem verimi ). Seçim yalnız kaynağı belgeler; "
                "η yukarıdaki 29. satırdan gelir.")
    g["A58"] = GIRIS_ACIKLAMA

    wb.save(yol)


def main():
    if not os.path.exists(AVAN):
        print("şablon bulunamadı:", AVAN)
        return 1
    yedek = AVAN + ".yedek"
    if not os.path.exists(yedek):
        shutil.copy2(AVAN, yedek)
        print("yedek alındı:", os.path.basename(yedek))
    yama(AVAN)
    print("yamalandı:", os.path.basename(AVAN))
    return 0


if __name__ == "__main__":
    sys.exit(main())

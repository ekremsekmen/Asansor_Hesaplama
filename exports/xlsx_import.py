# -*- coding: utf-8 -*-
"""
XLSX GERİ YÜKLEME  —  "Excel'den proje aç"

Programın ürettiği XLSX dosyası girdileri de taşır (girdi hücrelerinde,
proje antedi bilgisi ise dosya özelliklerinde).  Bu modül o dosyayı okuyup
girdileri arayüz alanlarına geri çevirir.

Kullanım senaryosu — REVİZYON:
    Proje klasöründeki Excel'i programa yükle → değişen girdiyi düzelt →
    güncel PDF ve XLSX'i yeniden indir.  Hiçbir şey baştan girilmez.

Hücre adresleri hucre_haritasi.py içindedir — dışa aktarma ile aynı harita.
Dosya Excel'de elle düzenlenmiş olsa bile girdiler oradan okunur.
"""
import io
import json

import openpyxl

from . import hucre_haritasi as H


class YuklemeHatasi(Exception):
    """Dosya bu programın ürettiği bir hesap dosyası değilse atılır."""


# ------------------------------------------------------------------ biçim
def _metin(x):
    """
    Hücre değerini arayüzün beklediği metne çevirir.
    Sayılarda Türkçe ondalık ayracı (virgül) kullanılır.
    """
    if x is None:
        return ""
    if isinstance(x, bool):
        return ""
    if isinstance(x, int):
        return str(x)
    if isinstance(x, float):
        if x != x or x in (float("inf"), float("-inf")):    # NaN / sonsuz
            return ""
        if float(x).is_integer():
            return str(int(x))
        return ("%.6f" % x).rstrip("0").rstrip(".").replace(".", ",")
    return str(x).strip()


def _oku(ws, adres, anahtar=None):
    """
    Hücreyi arayüz değerine çevirir.  Evet/Hayır kutuları METİN olarak
    saklandığı için mantıksal değere döndürülür — aksi hâlde "Hayır" metni
    işaret kutusunu doğru sayıp İŞARETLERDİ.
    """
    ham = ws[adres].value
    if anahtar and H.evet_hayir_mi(anahtar):
        return H.evet_mi(ham)
    return _metin(ham)


def _ek_nufus_oku(ws, tanim):
    satirlar = []
    for r in tanim["satirlar"]:
        kalem = _metin(ws[f"{tanim['kalem']}{r}"].value)
        miktar = _metin(ws[f"{tanim['miktar']}{r}"].value)
        if not kalem or miktar == "":
            continue
        satirlar.append({"aciklama": _metin(ws[f"{tanim['aciklama']}{r}"].value),
                         "miktar": miktar, "kalem": kalem})
    return satirlar


def _proje_oku(wb):
    """Proje antedi bilgisi dosya özelliklerinde saklanır."""
    p = {}
    ozellik = wb.properties
    ham = getattr(ozellik, "description", None)
    if ham:
        try:
            veri = json.loads(ham)
            if isinstance(veri, dict):
                p = {k: (veri.get(k) or "") for k in H.PROJE_ALANLARI}
        except (ValueError, TypeError):
            p = {}
    # JSON yoksa (ör. dosya başka bir programdan geçmişse) tek tek alanlardan topla
    if not p:
        p = {"proje_adi": ozellik.title or "", "isveren": ozellik.subject or "",
             "pafta_no": ozellik.category or "", "tarih": "",
             "muhendis": (ozellik.creator or "") if ozellik.creator != H.IMZA else ""}
    return {k: (p.get(k) or "") for k in H.PROJE_ALANLARI}


# ------------------------------------------------------------------ ana giriş
def xlsx_oku(icerik: bytes) -> dict:
    """
    Dönen yapı, arayüzün 'proje aç' işleviyle AYNI biçimdedir:
        { "tur": "tek" | "coklu" | "avan",
          "alanlar": { form_alani_kimligi: metin, ... },
          "proje":   { proje_adi, isveren, pafta_no, tarih, muhendis },
          "ozet":    kullanıcıya gösterilecek kısa açıklama }
    Böylece mevcut `uygula()` işlevi hiç değişmeden kullanılabilir.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(icerik), data_only=False)
    except Exception as e:                                    # noqa: BLE001
        raise YuklemeHatasi(
            "Dosya okunamadı — geçerli bir Excel (.xlsx) dosyası değil.") from e

    sayfalar = set(wb.sheetnames)
    alanlar, ek_nufus = {}, None

    # ---------------------------------------------------------- AVAN
    if H.AVAN_SAYFA in sayfalar and "ÖZET" in sayfalar:
        tur = "avan"
        ws = wb[H.AVAN_SAYFA]
        for anahtar, adres in H.AVAN_ORTAK.items():
            alanlar[H.avan_ortak_alan(anahtar)] = _oku(ws, adres, anahtar)
        kullanilan = 0
        for i, c in enumerate(H.AVAN_KOLONLAR, start=1):
            dolu = False
            for r, anahtar in H.AVAN_ASANSOR.items():
                deger = _oku(ws, f"{c}{r}", anahtar)
                alanlar[H.avan_asansor_alan(anahtar, i)] = deger
                if anahtar in ("kapasite", "Q_elle") and deger:
                    dolu = True
            alanlar[f"a_aktif{i}"] = dolu
            kullanilan += 1 if dolu else 0
        if H.AVAN_SABIT_SAYFA in sayfalar:
            wsS = wb[H.AVAN_SABIT_SAYFA]
            for anahtar, adres in H.AVAN_SABIT.items():
                alanlar[H.sabit_alan(anahtar)] = _oku(wsS, adres)
        if H.AVAN_AYD_SUTUN_SAYFA in sayfalar:
            alanlar[H.sabit_alan("ayd_sutun")] = _oku(
                wb[H.AVAN_AYD_SUTUN_SAYFA], H.AVAN_AYD_SUTUN_HUCRE)
        ozet = f"Avan hesapları — {kullanilan} asansör"

    # ---------------------------------------------------------- ÇOKLU TRAFİK
    elif H.COKLU_SAYFA in sayfalar:
        tur = "coklu"
        ws = wb[H.COKLU_SAYFA]
        for anahtar, adres in H.COKLU_ORTAK.items():
            alanlar[H.coklu_ortak_alan(anahtar)] = _oku(ws, adres)
        kullanilan = 0
        for i, c in enumerate(H.COKLU_KOLONLAR, start=1):
            for r, anahtar in H.COKLU_ASANSOR.items():
                deger = _oku(ws, f"{c}{r}")
                alanlar[H.coklu_asansor_alan(anahtar, i)] = deger
                if anahtar == "P" and deger:
                    kullanilan += 1
        ek_nufus = ("c", _ek_nufus_oku(ws, H.COKLU_EK_NUFUS))
        ozet = f"Çoklu asansör trafik hesabı — {kullanilan} asansör"

    # ---------------------------------------------------------- TEK TRAFİK
    elif H.TEK_SAYFA in sayfalar:
        tur = "tek"
        ws = wb[H.TEK_SAYFA]
        for anahtar, adres in H.TEK.items():
            alan = H.tek_alan(anahtar)
            if alan:                      # manuel_adet'in arayüzde karşılığı yok
                alanlar[alan] = _oku(ws, adres)
        ek_nufus = ("c", _ek_nufus_oku(ws, H.TEK_EK_NUFUS))
        ozet = "Tek asansör trafik hesabı"

    else:
        raise YuklemeHatasi(
            "Bu dosya programın ürettiği bir hesap dosyasına benzemiyor. "
            "Beklenen sayfalardan biri bulunamadı "
            f"({H.TEK_SAYFA} / {H.COKLU_SAYFA} / {H.AVAN_SAYFA}).")

    sonuc = {"tur": tur, "alanlar": alanlar, "proje": _proje_oku(wb), "ozet": ozet}
    if ek_nufus:
        sonuc["ek_nufus_hedef"], sonuc["ek_nufus"] = ek_nufus
    return sonuc

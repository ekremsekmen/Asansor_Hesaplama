# -*- coding: utf-8 -*-
"""
XLSX DIŞA AKTARIM  —  ŞABLON YÖNTEMİ

Program, sıfırdan bir tablo üretmez.  Ofisin kendi Excel dosyalarını ŞABLON
olarak açar, yalnız GİRDİ hücrelerini doldurur ve kaydeder.  Böylece:

  · tüm formüller, işlem adımları, sayfa düzeni ve pafta biçimi BİREBİR korunur,
  · dosya Excel'de açıldığında yeniden hesaplanır (fullCalcOnLoad),
  · ofis şablonu değişirse program çıktısı da kendiliğinden değişir,
  · üretilen dosya GİRDİLERİ de taşıdığı için programa geri yüklenebilir
    ( bkz. xlsx_import ) — revizyonda her şeyi baştan girmek gerekmez.

Hücre adresleri hucre_haritasi.py içindedir; dışa aktarma ve geri yükleme
aynı haritayı kullanır, bu yüzden ayrışamazlar.
"""
import io
import json
import os

import openpyxl

from engine.avan import hesap as E_AVAN
from engine.avan import trafik as E_TRF
from . import hucre_haritasi as H
from .sablon_denetim import dogrula

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates")
TRAFIK_SABLON = os.path.join(TEMPLATE_DIR, "ASANSOR_TRAFIK_HESABI_v2_1.xlsx")
AVAN_SABLON = os.path.join(TEMPLATE_DIR, "ASANSOR_AVAN_HESAPLARI.xlsx")


def _yaz(ws, adres, deger):
    """Boş/None değerleri hücreyi temizleyerek yazar (şablondaki örnek veriyi siler)."""
    #  Evet/Hayır kutuları Excel'de METİN olarak durur — şablon formülleri
    #  "Evet" karşılaştırması yapar; hücreye TRUE/FALSE yazmak bu kontrolü
    #  sessizce bozardı.
    if isinstance(deger, bool):
        deger = "Evet" if deger else "Hayır"
    ws[adres] = deger if deger not in ("", None) else None


def _temizle(ws, adresler):
    for ad in adresler:
        ws[ad] = None


def _proje_yaz(wb, proje):
    """
    Proje antedi bilgisini dosyanın ÖZELLİKLERİNE yazar.
    Şablonda bu bilgiye ait hücre olmadığı için ofis düzenine dokunulmaz;
    bilgi yine de dosyayla birlikte taşınır ve geri yüklenebilir.
    """
    p = {k: (proje or {}).get(k) for k in H.PROJE_ALANLARI}
    ozellik = wb.properties
    ozellik.title = p.get("proje_adi") or None
    ozellik.creator = p.get("muhendis") or H.IMZA
    ozellik.lastModifiedBy = p.get("muhendis") or H.IMZA
    ozellik.subject = p.get("isveren") or None
    ozellik.category = p.get("pafta_no") or None
    ozellik.keywords = H.IMZA
    ozellik.description = json.dumps(p, ensure_ascii=False)


def _kaydet(wb) -> bytes:
    wb.calculation.fullCalcOnLoad = True     # Excel açılışta yeniden hesaplasın
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# =====================================================================
#  TRAFİK HESABI  —  HESAPLAMA / PAFTA  ve  ÇOKLU ASANSÖR / PAFTA-COKLU
# =====================================================================
def trafik_xlsx(mod: str, g: dict, proje: dict = None) -> bytes:
    """mod: 'tek' | 'coklu'"""
    #  YANLIŞ / ESKİ ŞABLON KORUMASI:  şablon programın beklediği dosya değilse
    #  dosya ÜRETİLMEZ.  Sessizce yanlış bir pafta vermektense hiç vermemek
    #  doğrudur — hata mesajı hangi hücrenin uyuşmadığını söyler.
    dogrula(TRAFIK_SABLON, "trafik")
    wb = openpyxl.load_workbook(TRAFIK_SABLON)

    if mod == "tek":
        #  EKRAN NEYİ HESAPLADIYSA DOSYAYA DA O YAZILIR.  Arayüz girdileri
        #  `asansorler` listesinde gönderir; HESAPLAMA sayfası ise DÜZ alanları
        #  ( P / kapı / süreler / adet ) okur.  Bu eşlemeyi motorun kendisi
        #  yapar — burada ikinci bir kopya tutulursa ikisi ayrışır.
        g = E_TRF.tekil_girdi(g) or g
        ws = wb[H.TEK_SAYFA]
        for anahtar, adres in H.TEK.items():
            _yaz(ws, adres, g.get(anahtar))
        en = H.TEK_EK_NUFUS
        _temizle(ws, [f"{c}{r}" for r in en["satirlar"]
                      for c in (en["aciklama"], en["miktar"], en["kalem"])])
        for i, s in zip(en["satirlar"], (g.get("ek_nufus") or [])):
            _yaz(ws, f"{en['aciklama']}{i}", s.get("aciklama"))
            _yaz(ws, f"{en['miktar']}{i}", s.get("miktar"))
            _yaz(ws, f"{en['kalem']}{i}", s.get("kalem"))
    else:
        ws = wb[H.COKLU_SAYFA]
        for anahtar, adres in H.COKLU_ORTAK.items():
            _yaz(ws, adres, g.get(anahtar))
        _temizle(ws, [f"{c}{r}" for c in H.COKLU_KOLONLAR for r in H.COKLU_ASANSOR])
        for i, a in enumerate((g.get("asansorler") or [])[:4]):
            c = H.COKLU_KOLONLAR[i]
            for r, anahtar in H.COKLU_ASANSOR.items():
                _yaz(ws, f"{c}{r}", a.get(anahtar))
        en = H.COKLU_EK_NUFUS
        _temizle(ws, [f"{c}{r}" for r in en["satirlar"]
                      for c in (en["aciklama"], en["miktar"], en["kalem"])])
        for i, s in zip(en["satirlar"], (g.get("ek_nufus") or [])):
            _yaz(ws, f"{en['aciklama']}{i}", s.get("aciklama"))
            _yaz(ws, f"{en['miktar']}{i}", s.get("miktar"))
            _yaz(ws, f"{en['kalem']}{i}", s.get("kalem"))

    # ---- kullanılmayan hesap yolunu dosyadan çıkar
    #      Aksi hâlde "tek" indirmesinde PAFTA-COKLU sayfası şablonun kendi örnek
    #      verisiyle dolu, projeyle ilgisiz bir hesap gösterir; "çoklu" indirmesinde
    #      de PAFTA sayfası "HESAP HATASI" gösterir.  İki hesap yolu birbirinden
    #      bağımsız olduğu için (PAFTA yalnız HESAPLAMA'ya, PAFTA-COKLU yalnız
    #      ÇOKLU ASANSÖR'e bakar) ilgisiz olan güvenle silinebilir.
    if mod == "tek":
        hedef, silinecek, silinen_ad = "PAFTA", ["PAFTA-COKLU", H.COKLU_SAYFA], "PAFTA_COKLU_ALAN"
    else:
        hedef, silinecek, silinen_ad = "PAFTA-COKLU", ["PAFTA", H.TEK_SAYFA], "PAFTA_ALAN"
    for ad in silinecek:
        if ad in wb.sheetnames:
            del wb[ad]
    try:
        if silinen_ad in wb.defined_names:
            del wb.defined_names[silinen_ad]
    except Exception:
        pass

    # çıktı sayfası öne alınsın
    for sh in wb.worksheets:
        sh.views.sheetView[0].tabSelected = (sh.title == hedef)
    wb.active = wb.index(wb[hedef])
    _proje_yaz(wb, proje)
    return _kaydet(wb)


# =====================================================================
#  AVAN HESAPLARI  —  GİRİŞ / ÖZET / 1-4 NOLU ASANSÖR / MK.D. / TOPRAKLAMA
# =====================================================================
def avan_xlsx(veriler: dict, proje: dict = None) -> bytes:
    #  Ofis varsayılanları ve otomatik belirlenen değerler ( Nsç, L1 ) dosyaya
    #  AÇIKÇA yazılır: Excel'in kendi formülleri boş girdi hücresiyle çalışamaz,
    #  bu yüzden indirilen dosya ekrandakiyle aynı sonucu vermelidir.
    #  Burada yapılır ki hiçbir çağıran bu adımı atlayamasın.
    veriler = E_AVAN.girdileri_coz(veriler)
    dogrula(AVAN_SABLON, "avan")          # yanlış / eski şablon koruması
    wb = openpyxl.load_workbook(AVAN_SABLON)
    ws = wb[H.AVAN_SAYFA]
    o = veriler.get("ortak") or {}

    # 1) ORTAK GİRDİLER
    for anahtar, adres in H.AVAN_ORTAK.items():
        deger = o.get(anahtar)
        if anahtar in ("mk_uzunluk", "mk_genislik"):
            deger = deger or 0
        _yaz(ws, adres, deger)

    # 2) ASANSÖR BAZLI GİRDİLER  (C..F kolonları)
    _temizle(ws, [f"{c}{r}" for c in H.AVAN_KOLONLAR for r in H.AVAN_ASANSOR])
    for i, a in enumerate((veriler.get("asansorler") or [])[:4]):
        if not a:
            continue
        c = H.AVAN_KOLONLAR[i]
        for r, anahtar in H.AVAN_ASANSOR.items():
            _yaz(ws, f"{c}{r}", a.get(anahtar))

    # 3) SABİTLER sayfası — B bölümü (ofis standardı)
    sb = veriler.get("sabitler") or {}
    wsS = wb[H.AVAN_SABIT_SAYFA]
    for anahtar, adres in H.AVAN_SABIT.items():
        if sb.get(anahtar) is not None:
            _yaz(wsS, adres, sb[anahtar])
    if sb.get("ayd_sutun") is not None:
        _yaz(wb[H.AVAN_AYD_SUTUN_SAYFA], H.AVAN_AYD_SUTUN_HUCRE, sb["ayd_sutun"])

    # 4) MOTOR KORUMA CİHAZI — şablonda her asansör paftasında sabit metin
    #    ( "4 x 25" ) olarak duruyor ve motor gücünden bağımsızdı.  Program
    #    motor akımından seçtiği kademeyi buraya yazar; boş asansörün
    #    cetvelinde sigorta değeri kalmaz.
    hes = E_AVAN.hesapla(veriler)
    hesaplanan = {h.get("no"): h for h in (hes.get("asansorler") or []) if h}
    for i in range(1, 5):
        sayfa = H.avan_asansor_sayfasi(i)
        if sayfa not in wb.sheetnames:
            continue
        h = hesaplanan.get(i) or {}
        deger = (h.get("ozet") or {}).get("motor_sigorta") if h.get("aktif") else None
        _yaz(wb[sayfa], H.AVAN_SIGORTA_HUCRE, deger)

    for sh in wb.worksheets:
        sh.views.sheetView[0].tabSelected = (sh.title == "ÖZET")
    wb.active = wb.index(wb["ÖZET"])
    _proje_yaz(wb, proje)
    return _kaydet(wb)

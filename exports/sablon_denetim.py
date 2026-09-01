# -*- coding: utf-8 -*-
"""
ŞABLON DENETİMİ  —  "yanlış / eski Excel yüklenmiş mi?"

Program XLSX çıktısını ofisin kendi Excel dosyalarını ŞABLON olarak doldurarak
üretir.  Bu tasarımın tek riski şudur:  templates/ klasörüne YANLIŞ ya da ESKİ
bir dosya konursa program yine bir dosya üretir — ama pafta yanlış olur.
Hesap doğru olduğu için ekranda hiçbir belirti çıkmaz; hata yalnız teslim
edilen paftada görünür.  Bu modül o riski kapatır.

ÜÇ KATMANLI DENETİM
  1) YAPI    — beklenen sayfalar var mı?
  2) HÜCRE   — programın YAZDIĞI her girdi hücresi gerçekten yazılabilir mi?
               ( birleştirilmiş bir alanın ortasına düşmüş olmamalı — öyleyse
                 yazılan değer kaybolur )
  3) DEĞER   — şablonun İÇİNDEKİ tablolar motordaki tablolarla aynı mı?
               Eski bir şablonu yakalayan asıl kontrol budur:  ör. v1.3'te
               Tablo-4'e 1000/1200 mm satırları, Tablo-8'e 700 mm sütunu
               eklenmişti; eski dosyada bunlar yoktur.

Denetim BAŞARISIZSA XLSX üretilmez ( bkz. xlsx_export ) — sessizce yanlış
pafta vermektense hiç dosya vermemek doğrudur.
"""
import hashlib
import os

import openpyxl

from engine import tables as T

from . import hucre_haritasi as H


class SablonHatasi(Exception):
    """Şablon programın beklediği dosya değil."""


# =====================================================================
#  BEKLENEN SAYFALAR
# =====================================================================
TRAFIK_SAYFALARI = ("HESAPLAMA", "PAFTA", "ÇOKLU ASANSÖR", "PAFTA-COKLU",
                    "TABLO-1", "TABLO-2", "TABLO-4", "TABLO-6",
                    "TABLO-7", "TABLO-8", "TABLO-9", "TABLO-10")
AVAN_SAYFALARI = ("GİRİŞ", "ÖZET", "1 NOLU ASANSÖR", "2 NOLU ASANSÖR",
                  "3 NOLU ASANSÖR", "4 NOLU ASANSÖR", "MK.DAİRESİ AYD.",
                  "TOPRAKLAMA", "TABLOLAR", "SABİTLER")


# =====================================================================
#  YARDIMCILAR
# =====================================================================
def _sayi(x):
    """Hücre değerini sayıya çevirir; metin ('2,5' / '2.5') de kabul edilir."""
    if isinstance(x, bool) or x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    metin = str(x).strip().replace(",", ".")
    try:
        return float(metin)
    except ValueError:
        return None


def _esit(a, b, tolerans=1e-9):
    sa, sb = _sayi(a), _sayi(b)
    if sa is not None and sb is not None:
        return abs(sa - sb) <= tolerans
    return str(a).strip() == str(b).strip()


def _yazilabilir_mi(ws, adres):
    """
    Hücre BİRLEŞTİRİLMİŞ bir alanın ortasına düşüyorsa oraya yazılan değer
    Excel'de görünmez ve formüller onu okumaz — bu sessiz bir veri kaybıdır.
    """
    hucre = ws[adres]
    for aralik in ws.merged_cells.ranges:
        if hucre.coordinate in aralik and aralik.coord.split(":")[0] != hucre.coordinate:
            return False
    return True


def _girdi_hucreleri(tur):
    """Programın YAZDIĞI tüm hücreler: (sayfa, adres) listesi."""
    h = []
    if tur == "trafik":
        h += [(H.TEK_SAYFA, a) for a in H.TEK.values()]
        en = H.TEK_EK_NUFUS
        h += [(H.TEK_SAYFA, f"{en['kalem']}{en['satirlar'][0]}")]
        h += [(H.COKLU_SAYFA, a) for a in H.COKLU_ORTAK.values()]
        h += [(H.COKLU_SAYFA, f"{c}{r}")
              for c in H.COKLU_KOLONLAR for r in H.COKLU_ASANSOR]
        en = H.COKLU_EK_NUFUS
        h += [(H.COKLU_SAYFA, f"{en['kalem']}{en['satirlar'][0]}")]
    else:
        h += [(H.AVAN_SAYFA, a) for a in H.AVAN_ORTAK.values()]
        h += [(H.AVAN_SAYFA, f"{c}{r}")
              for c in H.AVAN_KOLONLAR for r in H.AVAN_ASANSOR]
        h += [(H.AVAN_SABIT_SAYFA, a) for a in H.AVAN_SABIT.values()]
        h += [(H.AVAN_AYD_SUTUN_SAYFA, H.AVAN_AYD_SUTUN_HUCRE)]
        h += [(H.avan_asansor_sayfasi(i), H.AVAN_SIGORTA_HUCRE) for i in range(1, 5)]
    return h


# =====================================================================
#  DEĞER DENETİMLERİ  —  şablondaki tablo  ↔  motordaki tablo
# =====================================================================
def _trafik_tablolari(wb, hata):
    #  TABLO-4  ( kapı açma / kapanma süreleri )  —  satır 4-10, A = genişlik,
    #  I..N = sayısal yardımcı kolonlar ( ta/tk × üç kapı tipi )
    ws = wb["TABLO-4"]
    tipler = ("Teleskopik Otomatik", "Merkezden Açılan Oto.", "Kabin İçi Oto. Kat K.Ç.")
    for i, genislik in enumerate(sorted(T.TABLO_4), start=4):
        okunan = ws[f"A{i}"].value
        if not _esit(okunan, genislik):
            hata.append(f"TABLO-4!A{i}: {genislik} mm bekleniyordu, {okunan!r} var "
                        "( şablon eski olabilir — 1000 / 1200 mm satırları v1.3'te eklendi )")
            continue
        for j, tip in enumerate(tipler):
            for k in range(2):                      # 0 = ta, 1 = tk
                sutun = chr(ord("I") + j * 2 + k)
                beklenen = T.TABLO_4[genislik][tip][k]
                okunan = ws[f"{sutun}{i}"].value
                if beklenen is None:
                    if str(okunan).strip().upper() not in ("YOK", "NONE", ""):
                        hata.append(f"TABLO-4!{sutun}{i}: 'YOK' bekleniyordu, {okunan!r} var")
                elif not _esit(okunan, beklenen):
                    hata.append(f"TABLO-4!{sutun}{i} ( {genislik} mm · {tip} ): "
                                f"{beklenen} bekleniyordu, {okunan!r} var")

    #  TABLO-8  ( kişi transfer süresi )  —  satır 1 genişlik, satır 2 tp
    ws = wb["TABLO-8"]
    for j, genislik in enumerate(sorted(T.TABLO_8)):
        sutun = chr(ord("B") + j)
        if not _esit(ws[f"{sutun}1"].value, genislik):
            hata.append(f"TABLO-8!{sutun}1: {genislik} mm bekleniyordu, "
                        f"{ws[f'{sutun}1'].value!r} var "
                        "( 700 mm sütunu v1.3'te eklendi )")
        elif not _esit(ws[f"{sutun}2"].value, T.TABLO_8[genislik]):
            hata.append(f"TABLO-8!{sutun}2 ( {genislik} mm ): "
                        f"{T.TABLO_8[genislik]} bekleniyordu, {ws[f'{sutun}2'].value!r} var")

    #  TABLO-7  ( kabin kapasitesi → anma yükü )  —  satır 2 kişi, satır 3 kg
    ws = wb["TABLO-7"]
    for j, kisi in enumerate(T.TABLO_7_BASILI):
        sutun = chr(ord("A") + j)
        if not _esit(_sayi(str(ws[f"{sutun}2"].value).split()[0]), kisi):
            hata.append(f"TABLO-7!{sutun}2: {kisi} kişi bekleniyordu, "
                        f"{ws[f'{sutun}2'].value!r} var")
        elif not _esit(_sayi(str(ws[f"{sutun}3"].value).split()[0]), T.TABLO_7[kisi]):
            hata.append(f"TABLO-7!{sutun}3 ( {kisi} kişi ): {T.TABLO_7[kisi]} kg "
                        f"bekleniyordu, {ws[f'{sutun}3'].value!r} var")

    #  TABLO-9  ( taşınacak insan yüzdesi )  —  D = standart, E = yükseltilmiş
    ws = wb["TABLO-9"]
    for i in range(2, 2 + len(T.TABLO_9)):
        ad = str(ws[f"A{i}"].value or "").strip()
        if ad not in T.TABLO_9:
            continue
        for sutun, anahtar in (("D", "Standart"), ("E", "Yükseltilmiş")):
            beklenen = T.TABLO_9[ad][anahtar]
            if not _esit(ws[f"{sutun}{i}"].value, beklenen):
                hata.append(f"TABLO-9!{sutun}{i} ( {ad} · {anahtar} ): {beklenen} "
                            f"bekleniyordu, {ws[f'{sutun}{i}'].value!r} var")

    #  TABLO-10  ( bekleme süresi sınırları )  —  F/G/H sayısal kolonlar
    ws = wb["TABLO-10"]
    for i in range(2, 2 + len(T.TABLO_10) + 2):
        ad = str(ws[f"A{i}"].value or "").strip()
        if ad not in T.TABLO_10:
            continue
        for sutun, anahtar in (("F", "sartli"), ("G", "standart"), ("H", "yukseltilmis")):
            beklenen = T.TABLO_10[ad][anahtar]
            if not _esit(ws[f"{sutun}{i}"].value, beklenen):
                hata.append(f"TABLO-10!{sutun}{i} ( {ad} · {anahtar} ): {beklenen} sn "
                            f"bekleniyordu, {ws[f'{sutun}{i}'].value!r} var")


def _avan_tablolari(wb, hata):
    ws = wb[H.AVAN_AYD_SUTUN_SAYFA]

    #  TABLO 2  —  oda aydınlatma verimi ızgarası ( satır 8-17 × sütun B-K )
    for i, k in enumerate(T.AYD_K_SATIRLARI):
        satir = 8 + i
        if not _esit(ws[f"A{satir}"].value, k):
            hata.append(f"TABLOLAR!A{satir}: aydınlatma k = {k} bekleniyordu, "
                        f"{ws[f'A{satir}'].value!r} var")
            continue
        for j, beklenen in enumerate(T.AYD_VERIM[i]):
            sutun = chr(ord("B") + j)
            if not _esit(ws[f"{sutun}{satir}"].value, beklenen):
                hata.append(f"TABLOLAR!{sutun}{satir} ( aydınlatma verimi ): {beklenen} "
                            f"bekleniyordu, {ws[f'{sutun}{satir}'].value!r} var")

    #  TABLO 7  —  kapasite / anma yükü  ( satır 41-49 )
    for i, kisi in enumerate(sorted(T.TABLO_7), start=41):
        if not _esit(ws[f"A{i}"].value, kisi):
            hata.append(f"TABLOLAR!A{i}: {kisi} kişi bekleniyordu, {ws[f'A{i}'].value!r} var")
        elif not _esit(ws[f"B{i}"].value, T.TABLO_7[kisi]):
            hata.append(f"TABLOLAR!B{i} ( {kisi} kişi ): {T.TABLO_7[kisi]} kg "
                        f"bekleniyordu, {ws[f'B{i}'].value!r} var")

    #  Kablo akım taşıma kapasitesi  ( satır 53-64 )
    for i, kesit in enumerate(sorted(T.KABLO_IZ), start=53):
        if not _esit(ws[f"A{i}"].value, kesit):
            hata.append(f"TABLOLAR!A{i}: {kesit} mm² bekleniyordu, {ws[f'A{i}'].value!r} var")
        elif not _esit(ws[f"B{i}"].value, T.KABLO_IZ[kesit]):
            hata.append(f"TABLOLAR!B{i} ( {kesit} mm² ): {T.KABLO_IZ[kesit]} A "
                        f"bekleniyordu, {ws[f'B{i}'].value!r} var")

    #  TABLO 11  —  anma yüküne göre boş kabin kütlesi  ( satır 70-78 )
    for i, (Q, Gk) in enumerate(T.TABLO_11, start=70):
        if not _esit(ws[f"A{i}"].value, Q):
            hata.append(f"TABLOLAR!A{i}: Q = {Q} kg bekleniyordu, {ws[f'A{i}'].value!r} var")
        elif not _esit(ws[f"B{i}"].value, Gk):
            hata.append(f"TABLOLAR!B{i} ( Q = {Q} kg ): Gk = {Gk} kg bekleniyordu, "
                        f"{ws[f'B{i}'].value!r} var")


# =====================================================================
#  ANA GİRİŞ
# =====================================================================
def _md5(yol):
    ozet = hashlib.md5()
    with open(yol, "rb") as f:
        for parca in iter(lambda: f.read(65536), b""):
            ozet.update(parca)
    return ozet.hexdigest()


#  Dosya başına önbellek — iki şablon dönüşümlü sorulduğunda birbirini
#  düşürmesin diye anahtar başına ayrı kayıt tutulur.
_ONBELLEK = {}
_ONBELLEK_AZAMI = 8


def _kaydet(anahtar, sonuc):
    if len(_ONBELLEK) >= _ONBELLEK_AZAMI:
        _ONBELLEK.clear()
    _ONBELLEK[anahtar] = sonuc


def denetle(yol: str, tur: str, onbellek: bool = True) -> dict:
    """
    tur: 'trafik' | 'avan'
    Dönen:
        {"uygun": bool, "hatalar": [...], "md5": "...", "dosya": "...",
         "sayfa_sayisi": n}
    Dosya değişmediği sürece sonuç önbellekten döner ( mtime + boyut ).
    """
    ad = os.path.basename(yol)
    if not os.path.exists(yol):
        return {"uygun": False, "dosya": ad, "md5": None, "sayfa_sayisi": 0,
                "hatalar": [f"Şablon dosyası bulunamadı: {ad}"]}

    st = os.stat(yol)
    anahtar = (yol, tur, st.st_mtime_ns, st.st_size)
    if onbellek and anahtar in _ONBELLEK:
        return _ONBELLEK[anahtar]

    hata = []
    try:
        wb = openpyxl.load_workbook(yol, data_only=False)
    except Exception as e:                                    # noqa: BLE001
        sonuc = {"uygun": False, "dosya": ad, "md5": None, "sayfa_sayisi": 0,
                 "hatalar": [f"Dosya açılamadı — geçerli bir .xlsx değil ({e})"]}
        _kaydet(anahtar, sonuc)
        return sonuc

    # 1) YAPI
    beklenen = TRAFIK_SAYFALARI if tur == "trafik" else AVAN_SAYFALARI
    eksik = [s for s in beklenen if s not in wb.sheetnames]
    if eksik:
        hata.append("Eksik sayfa: " + ", ".join(eksik))

    # 2) HÜCRE  —  yalnız sayfalar yerindeyse anlamlı
    if not eksik:
        for sayfa, adres in _girdi_hucreleri(tur):
            try:
                if not _yazilabilir_mi(wb[sayfa], adres):
                    hata.append(f"{sayfa}!{adres}: birleştirilmiş alanın ortasına "
                                "düşüyor — yazılan girdi kaybolur")
            except Exception:                                 # noqa: BLE001
                hata.append(f"{sayfa}!{adres}: hücre okunamadı")

        # 3) DEĞER
        try:
            (_trafik_tablolari if tur == "trafik" else _avan_tablolari)(wb, hata)
        except Exception as e:                                # noqa: BLE001
            hata.append(f"Tablo denetimi tamamlanamadı ({type(e).__name__}: {e}) — "
                        "şablonun düzeni beklenenden farklı")

    sonuc = {"uygun": not hata, "dosya": ad, "md5": _md5(yol),
             "sayfa_sayisi": len(wb.sheetnames), "hatalar": hata}
    _kaydet(anahtar, sonuc)
    return sonuc


def dogrula(yol: str, tur: str):
    """Uygun değilse SablonHatasi atar — XLSX üretim yolunda kullanılır."""
    s = denetle(yol, tur)
    if s["uygun"]:
        return s
    n = len(s["hatalar"])
    ilk = "  ·  ".join(s["hatalar"][:3])
    if n > 3:
        ilk += f"  ·  ( ve {n - 3} sorun daha )"
    raise SablonHatasi(
        f"ŞABLON UYUŞMUYOR — {s['dosya']} bu programın beklediği dosya değil, "
        f"XLSX üretilmedi. {ilk}   Doğru şablonu templates/ klasörüne koyun; "
        "Sabitler sekmesindeki 'Şablon durumu' bölümü ayrıntıyı gösterir.")

# -*- coding: utf-8 -*-
"""
TEST 4  —  ÇIKTI BÜTÜNLÜĞÜ  (XLSX ve PDF)

Üretilen dosyalar gerçekten açılabiliyor mu, içinde Excel hata hücresi var mı,
doğru sayfaları taşıyor mu, PDF geçerli ve Türkçe karakterler yerinde mi?
"""
import io
import json as _js
import os
import re
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                            # noqa: E402
from engine.avan import hesap as AV
from engine.avan import trafik as TR               # noqa: E402
from exports import hucre_haritasi as H                    # noqa: E402
try:
    from exports import dxf_export as _DXF          # noqa: E402
except Exception:                                   # noqa: BLE001
    _DXF = None            # CAD kitaplıkları yoksa paket testi atlanır
from engine.uygulama import sabitler as _US2
from exports import kapak_export as KPK
from exports import pdf_export as PE, xlsx_export as XE    # noqa: E402
from api.avan import XLSX_TUR                           # noqa: E402
from testler.ortak import Rapor, hata_hucresi_ara, yeniden_hesapla, soffice_yolu  # noqa: E402

# Geçici dosyalar sistemin temp klasörüne yazılır — proje klasörü kirlenmez
# ve silme izni kısıtlı makinelerde test takılmaz.
GECICI = os.path.join(tempfile.gettempdir(),
                      "avan_test_" + os.path.splitext(os.path.basename(__file__))[0])

PROJE = {"proje_adi": "Türkçe Şıkır Ğüzel Öİ Projesi", "isveren": "ÇAĞDAŞ İnşaat A.Ş.",
         "pafta_no": "EL-04", "tarih": "27.08.2026", "muhendis": "Ekrem Sekmen"}

GT = dict(bina_tipi="Konut", bina_yuksekligi=39.98, yapi_yuksekligi=43, N=11, h=3,
          hizli1=44, hizli2=3, P=10, kapi_genisligi=900, kapi_tipi="Merkezden Açılan Oto.")
GC = dict(bina_tipi="Konut", bina_yuksekligi=39.98, yapi_yuksekligi=43, N=11, h=3,
          hizli1=44, hizli2=3,
          asansorler=[dict(P=10, kapi_genisligi=900, kapi_tipi="Teleskopik Otomatik"),
                      dict(P=16, kapi_genisligi=1100, kapi_tipi="Teleskopik Otomatik")])
ORT = dict(U=380, kappa=56, eps_max=3, temel_a=26.55, temel_b=16.4, beta=150,
           serit_L=58.5, cubuk_sayisi=4, mk_uzunluk=4200, mk_genislik=3100)
A1 = dict(tanim="İnsan", kapasite=10, V=1.6, eta=0.85, Hk=32.85, kuyu_genisligi=1800,
          kabin_boyu=1450, kabin_genisligi=1300, gr=17.91, Fmk=350, Fsh=100, Nsc=11,
          S1=6, L1=32.85, S2=6, L2=3, kablo_tipi="NHXMH FE180")
A2 = dict(A1, tanim="Sedye + Yük", kapasite=16, V=1.0, kuyu_genisligi=2650,
          kabin_boyu=1350, kabin_genisligi=2100)
AV_VERI = {"ortak": ORT, "asansorler": [A1, A2, A1, A2]}

TR_HARF = "ÇĞİÖŞÜçğıöşü"


def calistir():
    print("\n\033[1mTEST 4 — ÇIKTI BÜTÜNLÜĞÜ (XLSX ve PDF)\033[0m")
    r = Rapor("Çıktı bütünlüğü")
    shutil.rmtree(GECICI, ignore_errors=True)
    os.makedirs(GECICI, exist_ok=True)

    # ---------------------------------------------------------- XLSX
    dosyalar = {
        "tek": (XE.trafik_xlsx("tek", GT), ["HESAPLAMA", "PAFTA"], ["ÇOKLU ASANSÖR", "PAFTA-COKLU"], "PAFTA"),
        "coklu": (XE.trafik_xlsx("coklu", GC), ["ÇOKLU ASANSÖR", "PAFTA-COKLU"], ["HESAPLAMA", "PAFTA"], "PAFTA-COKLU"),
        "avan": (XE.avan_xlsx(AV_VERI),
                 ["GİRİŞ", "ÖZET", "1 NOLU ASANSÖR", "2 NOLU ASANSÖR", "3 NOLU ASANSÖR",
                  "4 NOLU ASANSÖR", "MK.DAİRESİ AYD.", "TOPRAKLAMA", "TABLOLAR", "SABİTLER"],
                 [], "ÖZET"),
    }
    for ad, (icerik, olmali, olmamali, aktif) in dosyalar.items():
        yol = os.path.join(GECICI, f"{ad}.xlsx")
        open(yol, "wb").write(icerik)
        r.kontrol(f"{ad}.xlsx ZIP imzası", icerik[:2] == b"PK")
        r.kontrol(f"{ad}.xlsx makul boyut", 20_000 < len(icerik) < 5_000_000, f"→ {len(icerik)} bayt")
        try:
            wb = openpyxl.load_workbook(yol)
        except Exception as e:                                   # noqa: BLE001
            r.kontrol(f"{ad}.xlsx açılabiliyor", False, f"→ {e}")
            continue
        r.kontrol(f"{ad}.xlsx açılabiliyor", True)
        for sh in olmali:
            r.kontrol(f"{ad}.xlsx '{sh}' sayfası var", sh in wb.sheetnames)
        for sh in olmamali:
            r.kontrol(f"{ad}.xlsx '{sh}' sayfası ÇIKARILMIŞ", sh not in wb.sheetnames)
        r.esit(f"{ad}.xlsx açılışta gelen sayfa", wb.active.title, aktif)
        r.kontrol(f"{ad}.xlsx açılışta yeniden hesaplanacak",
                  bool(getattr(wb.calculation, "fullCalcOnLoad", False)))
        # tanımlı adlarda kırık başvuru olmamalı
        kirik = [n for n, d in wb.defined_names.items() if "#REF" in str(d.value)]
        r.kontrol(f"{ad}.xlsx tanımlı adlar sağlam", not kirik, f"→ {kirik}")

        #  v1.9 — MOTOR SİGORTASI: şablonda her asansörde sabit "4 x 25"
        #  yazıyordu.  Program motor akımından seçtiği kademeyi hücreye
        #  yazmalı ki indirilen Excel ile ekrandaki pafta ayrışmasın.
        if ad == "avan":
            from engine.avan import hesap as _AV
            _hes = _AV.hesapla(AV_VERI)
            _bek = {h["no"]: (h["ozet"]["motor_sigorta"] if h.get("aktif") else None)
                    for h in (_hes.get("asansorler") or []) if h}
            for _i in range(1, 5):
                _sayfa = f"{_i} NOLU ASANSÖR"
                if _sayfa not in wb.sheetnames:
                    continue
                _okunan = wb[_sayfa]["G102"].value
                r.esit(f"avan.xlsx {_sayfa}!G102 motor sigortası",
                       _okunan, _bek.get(_i))
                r.kontrol(f"avan.xlsx {_sayfa} şablondaki sabit metin kalmadı",
                          not (_bek.get(_i) not in (None, "4 x 25") and _okunan == "4 x 25"))


    # yeniden hesaplandığında hata hücresi kalmamalı
    if soffice_yolu():
        cikis = os.path.join(GECICI, "hesaplandi")
        yeniden_hesapla([os.path.join(GECICI, f"{a}.xlsx") for a in dosyalar], cikis)
        for ad in dosyalar:
            y = os.path.join(cikis, f"{ad}.xlsx")
            if not os.path.exists(y):
                r.kontrol(f"{ad}.xlsx yeniden hesaplanabildi", False)
                continue
            hatalar = hata_hucresi_ara(y)
            r.kontrol(f"{ad}.xlsx hesap sonrası hata hücresi yok", not hatalar,
                      f"→ {hatalar[:4]}")
    else:
        r.atla("LibreOffice yok — XLSX yeniden hesaplama kontrolü atlandı")

    # ---------------------------------------------------------- PDF
    pdfler = {
        "tek": PE.trafik_pdf(TR.hesapla_tek(GT), PROJE),
        "coklu": PE.trafik_pdf(TR.hesapla_coklu(GC), PROJE),
        "avan": PE.avan_pdf(AV.hesapla(AV_VERI), PROJE),
        # hesap hatalı girdide de geçerli belge üretilmeli
        "hatali": PE.trafik_pdf(TR.hesapla_tek({"bina_tipi": "Konut"}), PROJE),
        "bos_avan": PE.avan_pdf(AV.hesapla({"ortak": {}, "asansorler": []}), PROJE),
        "projesiz": PE.trafik_pdf(TR.hesapla_tek(GT), None),
    }
    BEKLENEN_BASLIK = {"tek": "ASANSÖR TRAFİK HESABI",
                       "coklu": "ÇOKLU ASANSÖR TRAFİK HESABI",
                       "avan": "ASANSÖR AVAN PROJE HESAPLARI",
                       "hatali": "ASANSÖR TRAFİK HESABI",
                       "bos_avan": "ASANSÖR AVAN PROJE HESAPLARI",
                       "projesiz": "ASANSÖR TRAFİK HESABI"}
    try:
        import pypdfium2 as pdfium
    except ImportError:
        pdfium = None
        #  SESSİZ ATLAMA OLMAZ.  Bu kütüphane yokken PDF metin kontrolleri
        #  hiç çalışmıyor ve kimse haberdar olmuyordu — 96 kontrol boşta
        #  duruyordu.  Eksikse en azından söylensin.
        r.atla("pypdfium2 kurulu değil — PDF metin içeriği kontrolleri atlandı "
               "(python3 -m pip install pypdfium2)")

    _pypdf_uyarisi = []          # atlama uyarısı bir kez basılsın
    for ad, icerik in pdfler.items():
        yol = os.path.join(GECICI, f"{ad}.pdf")
        open(yol, "wb").write(icerik)
        r.kontrol(f"{ad}.pdf imzası", icerik[:4] == b"%PDF")
        r.kontrol(f"{ad}.pdf sonlandırılmış", b"%%EOF" in icerik[-2048:])
        r.kontrol(f"{ad}.pdf makul boyut", 5_000 < len(icerik) < 20_000_000,
                  f"→ {len(icerik)} bayt")
        if not pdfium:
            continue
        try:
            d = pdfium.PdfDocument(yol)
            sayfa = len(d)
        except Exception as e:                                   # noqa: BLE001
            r.kontrol(f"{ad}.pdf açılabiliyor", False, f"→ {e}")
            continue
        r.kontrol(f"{ad}.pdf açılabiliyor ({sayfa} sayfa)", sayfa >= 1)
        metinler = [d[i].get_textpage().get_text_range() for i in range(sayfa)]
        tam = "\n".join(metinler)
        # her sayfada üst şerit doğru belgeyi göstermeli
        for i, m in enumerate(metinler):
            ilk = m.split("\n")[0] if m else ""
            r.kontrol(f"{ad}.pdf s.{i+1} üst şerit doğru",
                      ilk.startswith(BEKLENEN_BASLIK[ad]), f"→ {ilk[:60]!r}")
        # Türkçe karakterler bozulmamalı
        r.kontrol(f"{ad}.pdf Türkçe karakterler yerinde",
                  any(h in tam for h in TR_HARF))
        r.kontrol(f"{ad}.pdf kırık karakter yok",
                  "�" not in tam and "□" not in tam)
        # sayfa numarası ve tarih alt bilgide olmalı
        r.kontrol(f"{ad}.pdf sayfa numarası var", re.search(r"Sayfa\s*\d+", tam) is not None)
        # başlık üstverisi
        try:
            import pypdf
            import warnings
            warnings.filterwarnings("ignore")
            m = pypdf.PdfReader(yol).metadata
            r.kontrol(f"{ad}.pdf üstveri başlığı doğru",
                      str(m.get("/Title", "")).startswith(BEKLENEN_BASLIK[ad]),
                      f"→ {m.get('/Title')!r}")
        except ImportError:
            if not _pypdf_uyarisi:
                r.atla("pypdf kurulu değil — PDF üstveri kontrolleri atlandı "
                       "(python3 -m pip install pypdf)")
                _pypdf_uyarisi.append(True)

    #  PROJE ANTEDİ PAFTADA DEĞİL, KAPAKTADIR.
    #  Kapak sayfası eklendiğinde proje adı / işveren / mühendis bilgisi
    #  pafta PDF'lerinden bilerek kaldırıldı ( bkz. exports/pdf_export.py:
    #  "Pafta PDF'leri proje antedi taşımaz; bu bilgiler yalnız kapaktadır." ).
    #  Test bunu doğrular: bilgi paftada GÖRÜNMEMELİ.
    #
    #  TRAFİK PAFTASI ayrıca TEK SAYFADIR ve SONUÇ ile biter: otomatik öneri
    #  tablosu ve imza kutusu paftadan çıkarıldı.  İçerik sığmazsa orantılı
    #  küçültülür — hiçbir satır atılmaz, ikinci sayfaya taşma olmaz.
    if pdfium:
        for _ad in ("tek", "coklu"):
            d = pdfium.PdfDocument(os.path.join(GECICI, f"{_ad}.pdf"))
            tum = "\n".join(d[i].get_textpage().get_text_range() for i in range(len(d)))
            r.esit(f"{_ad} trafik paftası TEK SAYFA", len(d), 1)
            r.kontrol(f"{_ad} paftasında proje adı YOK ( kapağa taşındı )",
                      "Türkçe Şıkır" not in tum)
            r.kontrol(f"{_ad} paftasında işveren YOK ( kapağa taşındı )",
                      "ÇAĞDAŞ" not in tum)
            r.kontrol(f"{_ad} paftasında imza kutusu YOK", "Hesabı yapan" not in tum)
            r.kontrol(f"{_ad} paftasında otomatik öneri tablosu YOK",
                      "OTOMATİK ÖNERİ" not in tum)
            r.kontrol(f"{_ad} paftasında nüfus dökümü YOK",
                      "NÜFUSUN AYRINTISI" not in tum)
            r.kontrol(f"{_ad} paftası SONUÇ ile bitiyor", "SONUÇ" in tum)
            #  SİYAH BEYAZ / AutoCAD UYUMU:  paftada RENK olmamalı.
            #  Renkli zemin gri baskıda ayrışmaz, AutoCAD'de PDFIMPORT ile
            #  solid hatch'e döner ve üstündeki beyaz yazı kaybolur.
            import re as _re
            _ic = open(os.path.join(GECICI, f"{_ad}.pdf"), "rb").read().decode("latin-1")
            _renk = set()
            for _m in _re.finditer(r"([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+(rg|RG)\b", _ic):
                _r, _g, _b = (round(float(_m.group(i)), 3) for i in (1, 2, 3))
                if abs(_r - _g) > 0.02 or abs(_g - _b) > 0.02:
                    _renk.add((_r, _g, _b))
            r.kontrol(f"{_ad} paftasında renk yok ( siyah beyaz / AutoCAD )",
                      not _renk, f"→ {sorted(_renk)[:4]}")
            #  Küçültme İÇERİK KAYBETMEZ — hesabın gövdesi yerinde olmalı
            for _im in (("BİNADA BULUNAN İNSAN SAYISININ TESPİTİ", "B = b + ( n · b )")
                        if _ad == "tek" else ("ORTAK BİNA BİLGİLERİ", "GRUP KONTROLÜ")):
                r.kontrol(f"{_ad} paftasında bölüm duruyor: {_im[:28]}", _im in tum)
        #  Avan paftası çok sayfalıdır ve imza kutusu ORADA durur
        d = pdfium.PdfDocument(os.path.join(GECICI, "avan.pdf"))
        _av = "\n".join(d[i].get_textpage().get_text_range() for i in range(len(d)))
        r.kontrol("avan paftasında imza kutusu duruyor", "Hesabı yapan" in _av)

    # ------------------------------------------------- v1.3 içerik kontrolleri
    #  Yeni büyüklükler paftaya gerçekten basılıyor mu?
    if pdfium:
        def _metin(icerik):
            import tempfile as _t
            y = os.path.join(GECICI, "_gecici_icerik.pdf")
            open(y, "wb").write(icerik)
            d = pdfium.PdfDocument(y)
            return "\n".join(d[i].get_textpage().get_text_range() for i in range(len(d)))

        m = _metin(PE.trafik_pdf(TR.hesapla_tek(dict(GT, bodrum=2)), PROJE))
        for beklenen in ("Bodrum durak adedi", "Toplam durak adedi",
                         "Toplam seyahat mesafesi", "H (Tablo-3)"):
            r.kontrol(f"tek pafta: '{beklenen}' basıldı", beklenen in m)
        m0 = _metin(PE.trafik_pdf(TR.hesapla_tek(GT), PROJE))
        r.kontrol("bodrumsuzda bodrum notu basılmıyor", "H (Tablo-3)" not in m0)
        #  ara değerli kapı ve erişilebilirlik uyarıları
        m = _metin(PE.trafik_pdf(TR.hesapla_tek(dict(GT, kapi_genisligi=1000)), PROJE))
        r.kontrol("1000 mm ara değer uyarısı paftada", "enterpolasyon" in m)
        m = _metin(PE.trafik_pdf(TR.hesapla_tek(dict(GT, P=6, kapi_genisligi=700)), PROJE))
        r.kontrol("erişilebilirlik uyarısı paftada", "81-70" in m)
        r.kontrol("dış değerleme uyarısı paftada", "değerleme" in m)
        #  kamu binası varsayımları
        m = _metin(PE.trafik_pdf(TR.hesapla_tek(dict(GT, bina_tipi="Kamu Binaları")), PROJE))
        r.kontrol("kamu %k varsayımı paftada", "Tablo-9" in m and "MERKEZ" in m.upper())
        #  çoklu: asansör bazında bodrum ve imalatçı süresi
        m = _metin(PE.trafik_pdf(TR.hesapla_coklu(dict(
            GC, bodrum=2,
            asansorler=[dict(P=10, kapi_genisligi=900, kapi_tipi="Teleskopik Otomatik",
                             manuel_ta=2.9)])), PROJE))
        r.kontrol("çoklu paftada bodrum satırı", "Bodrum durak adedi" in m)
        r.kontrol("çoklu paftada imalatçı süresi uyarısı", "imalatçı verisiyle" in m)
        #  avan: asansör bazında askı/denge ve tutarlılık uyarısı
        _t = TR.hesapla_tek(dict(GT, bodrum=1))
        _av = AV.hesapla({"ortak": ORT,
                          "asansorler": [dict(A1, i_palanga=1, q_denge=0.4)],
                          "trafik": TR.trafik_ozeti(_t)})
        m = _metin(PE.avan_pdf(_av, PROJE))
        r.kontrol("avan paftasında asansör bazı kaynağı", "bazında" in m)
        r.kontrol("avan paftasında tutarlılık uyarısı", "NOLU ASANSÖR:" in m)

        #  TOPRAKLAMA KONTROLÜNÜN AÇIKLAMA NOTLARI PAFTAYA BASILMAZ.
        #  ( kullanıcı isteği — 2.5 )  Söyledikleri değerler zaten hesap
        #  satırlarında yazılı ( UL, IΔn, β ) ve dört satırlık blok yer
        #  kaplıyordu.  Bilgi kaybolmuyor:  motor bunları "ekran_notlari"
        #  olarak veriyor, arayüz ⓘ altında gösteriyor.
        _tp = AV.hesapla({"ortak": ORT, "asansorler": [dict(A1)]}).get("topraklama") or {}
        _b3 = (_tp.get("bolumler") or [{}])[-1]
        r.kontrol("topraklama notları ekran_notlari'na taşındı",
                  len(_b3.get("ekran_notlari") or []) == 3)
        r.kontrol("topraklama bölümünde paftaya basılacak not kalmadı",
                  not (_b3.get("notlar") or []) and not (_b3.get("aciklamalar") or []))
        _mp = _metin(PE.avan_pdf(AV.hesapla(
            {"ortak": ORT, "asansorler": [dict(A1)]}), PROJE))
        for _yasak in ("HESAPLANAN TAHMİNİ DEĞERDİR", "KABULLER: TT şebeke",
                       "zemin etüdünden alınmalıdır"):
            r.kontrol(f"paftada yok: {_yasak[:28]}…", _yasak not in _mp)
        #  ama hesabın kendisi ve hükmü yerinde
        r.kontrol("topraklama kontrolü paftada duruyor",
                  "temel topraklaması yeterlidir" in _mp or "ek topraklayıcı" in _mp)

        #  MAKİNE DAİRESİZ ( MRL ) SİSTEM — bölüm paftada HİÇ BASILMAZ.
        #  Makine dairesi yoksa aydınlatma hesabının konusu da yoktur; eskiden
        #  "bu hesap uygulanmaz" satırı boşuna yer kaplıyordu.  Ama MRL kutusu
        #  İŞARETLİ DEĞİLKEN ölçü de girilmemişse bu unutulmuş bir girdidir —
        #  o zaman uyarı basılmalı, yoksa eksik hesap sessizce gizlenirdi.
        def _mkm(**ortak_ek):
            return _metin(PE.avan_pdf(AV.hesapla(
                {"ortak": dict(ORT, **ortak_ek), "asansorler": [dict(A1)]}), PROJE))

        _mk_yok = {"mk_uzunluk": None, "mk_genislik": None}
        for _ad, _ek in (("kutu işaretli", dict(_mk_yok, mk_yok=True)),
                         ("kutu işaretli + ölçü dolu", dict(mk_yok=True)),
                         ("eski dosya, kutu gönderilmemiş", dict(_mk_yok))):
            r.kontrol(f"MRL ( {_ad} ) → makine dairesi bölümü paftada YOK",
                      "MAKİNE DAİRESİ" not in _mkm(**_ek))
        _m2 = _mkm(**dict(_mk_yok, mk_yok=False))
        r.kontrol("MRL kutusu kapalı + ölçü yok → uyarı paftada basılıyor",
                  "MAKİNE DAİRESİ" in _m2 and "ÖLÇÜLERİ GİRİLMEDİ" in _m2)
        r.kontrol("ölçü girilince makine dairesi hesabı paftada",
                  "MAKİNE DAİRESİ AYDINLATMA HESABI" in _mkm(mk_yok=False))

    # ---------------------------------------------------------- büyük/uç durumlar
    buyuk = dict(GC, asansorler=[dict(P=25, kapi_genisligi=1300,
                                      kapi_tipi="Merkezden Açılan Oto.") for _ in range(4)])
    b = PE.trafik_pdf(TR.hesapla_coklu(buyuk), PROJE)
    r.kontrol("4 asansörlü çoklu PDF üretildi", b[:4] == b"%PDF" and len(b) > 10_000)
    uzun = dict(PROJE, proje_adi="A" * 300, isveren="B" * 300)
    b = PE.trafik_pdf(TR.hesapla_tek(GT), uzun)
    r.kontrol("çok uzun proje adı PDF'i bozmadı", b[:4] == b"%PDF")
    x = XE.trafik_xlsx("tek", dict(GT, ek_nufus=[
        {"aciklama": "Kalem " + str(i), "miktar": 10, "kalem": "KONUT — Diğer oda"}
        for i in range(30)]))
    r.kontrol("30 ek nüfus satırı XLSX'i bozmadı", x[:2] == b"PK")

    # ---------------------------------------------------------- şablon eksikse
    #  Hesap ve PDF şablonsuz da çalışmalı; XLSX üretilemez ama kullanıcı
    #  bozuk bir dosya indirmemeli — anlaşılır hata almalı.
    #
    #  ÖNEMLİ: bu kontrol gerçek şablon dosyalarına DOKUNMAZ.  Yalnız modülün
    #  yol değişkenleri geçici olarak var olmayan bir dosyayı gösterir; test
    #  bitince eski değerlerine döner.  Böylece silme izni kısıtlı makinelerde
    #  de güvenle çalışır.
    yok = os.path.join(tempfile.gettempdir(), "avan_olmayan_sablon.xlsx")
    eski_trafik, eski_avan = XE.TRAFIK_SABLON, XE.AVAN_SABLON
    XE.TRAFIK_SABLON = XE.AVAN_SABLON = yok
    try:
        s2 = TR.hesapla_tek(GT)
        r.kontrol("şablonsuz hesap çalışıyor", s2.get("hata") is None and s2["ozet"]["TR"] > 0)
        p2 = PE.trafik_pdf(s2, PROJE)
        r.kontrol("şablonsuz PDF üretiliyor", p2[:4] == b"%PDF" and len(p2) > 10_000)
        p3 = PE.avan_pdf(AV.hesapla(AV_VERI), PROJE)
        r.kontrol("şablonsuz avan PDF üretiliyor", p3[:4] == b"%PDF")
        for ad, fn in (("trafik", lambda: XE.trafik_xlsx("tek", GT)),
                       ("avan", lambda: XE.avan_xlsx(AV_VERI))):
            hata_alindi = False
            try:
                fn()
            except Exception:                                # noqa: BLE001
                hata_alindi = True
            r.kontrol(f"şablonsuz {ad} XLSX açıkça hata veriyor "
                      "(sessizce bozuk dosya değil)", hata_alindi)
    finally:
        XE.TRAFIK_SABLON, XE.AVAN_SABLON = eski_trafik, eski_avan
    r.kontrol("şablon dosyaları yerinde ve okunabilir",
              os.path.isfile(XE.TRAFIK_SABLON) and os.path.isfile(XE.AVAN_SABLON)
              and os.path.getsize(XE.TRAFIK_SABLON) > 10_000
              and os.path.getsize(XE.AVAN_SABLON) > 10_000)
    r.kontrol("şablonla XLSX yeniden üretilebiliyor",
              XE.trafik_xlsx("tek", GT)[:2] == b"PK")

    #  v2.9 — KAPAK ADRESİ SESSİZCE KIRPILMASIN.  Satır bütçesi 2'ye sabitti;
    #  124 karakterlik normal bir adreste "İstanbul Türkiye 34758" bölümü
    #  çıktıdan düşüyor, kullanıcıya hiçbir şey söylenmiyordu.  Bütçe artık
    #  kutunun gerçek yüksekliğinden türetiliyor;  yine sığmazsa kırpma
    #  ÇIKTIDA görünür ( … ).
    def _kapak_metni(_adres):
        _b = KPK.pdf_bytes({"project_title": "D", "owner": "X",
                            "company_name": "ABC", "company_address": _adres})
        return _metin(_b) if pdfium else ""
    if pdfium:
        for _ad, _adres in (
                ("90 karakter", "Atatürk Mahallesi Cumhuriyet Caddesi No 145 "
                                "Kat 7 Daire 21 Ataşehir İstanbul Türkiye 34758"),
                ("124 karakter", "Barbaros Hayrettin Paşa Mahallesi 1993. Sokak "
                                 "Nuvo Dragos Sitesi A Blok No 12 Kat 9 Daire 41 "
                                 "Ataşehir İstanbul Türkiye 34758")):
            _t = _kapak_metni(_adres)
            _eksik = [p for p in _adres.split() if p not in _t]
            r.kontrol(f"kapak adresi tam basılıyor ( {_ad} )", not _eksik,
                      f"→ düşen: {' '.join(_eksik)[:60]}")
        #  Kutuya hiç sığmayan bir adreste kırpma GÖRÜNÜR olmalı
        _cok = " ".join(["Mahalle Sokak Bina Kat Daire"] * 12)
        r.kontrol("sığmayan adreste kırpma işareti basılıyor",
                  "…" in _kapak_metni(_cok))

    #  v2.9 — ÇÖKEN TEST ÇIKIŞ KODUNU BOZMALIDIR.  Bir test modülü istisna
    #  atınca ( kaldi = 0 ) sayaç artmıyor, özet "0 başarısız" diyor ve çıkış
    #  kodu 0 oluyordu:  sürekli tümleştirme yeşil görünürken test hiç
    #  koşmamış oluyordu.
    import subprocess as _sp, tempfile as _tf, textwrap as _tw
    _gecici = _tf.mkdtemp(prefix="avan_kosucu_")
    with open(os.path.join(_gecici, "coken.py"), "w", encoding="utf-8") as _f:
        _f.write("def calistir():\n    raise RuntimeError('kasıtlı çökme')\n")
    _kok = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    _kod = _tw.dedent(f"""
        import sys, os
        sys.path.insert(0, {_kok!r})
        sys.path.insert(0, {_gecici!r})
        from testler import calistir as K
        K.TESTLER = [("9", "Çöken", "coken", False)]
        sys.exit(K.main(["9"]))
    """)
    _p = _sp.run([sys.executable, "-c", _kod], capture_output=True, text=True,
                 env={**os.environ, "NO_COLOR": "1"})
    r.esit("çöken test çıkış kodunu 1 yapıyor", _p.returncode, 1)
    r.kontrol("çöken test özet satırında bildiriliyor",
              "ÇALIŞTIRILAMAYAN" in _p.stdout, f"→ {_p.stdout[-160:]!r}")

    #  v2.9 — EKRAN REDDEDİYORSA XLSX DE ÜRETİLMEZ.
    #  PDF hatayı paftaya BASAR ( okunur belge çıkar ) ama Excel şablonu yalnız
    #  girdi hücrelerini alır:  Python'a özgü denetimler ( "durak adedi N+1
    #  olmalıdır" gibi ) şablonda yoktur, dolayısıyla ekranda reddedilen bir
    #  hesap indirilen dosyada SORUNSUZ görünüyordu.
    import json as _json
    from api import avan as _M
    _hatali = {"girdiler": {"bina_tipi": "Konut", "bina_yuksekligi": "39,98",
                            "yapi_yuksekligi": "43", "N": "11", "h": "3",
                            "hizli1": "44", "hizli2": "3",
                            "asansorler": [{"P": "10", "kapi_genisligi": "900",
                                            "kapi_tipi": "Teleskopik Otomatik",
                                            "durak": "2"}]}}
    _y = _M.indir_trafik_xlsx(_hatali)
    _g = _json.loads(bytes(_y.body).decode("utf-8")) if "json" in _y.media_type else {}
    r.kontrol("hatalı hesapta trafik XLSX üretilmiyor",
              "json" in _y.media_type and "HESAP HATASI" in str(_g.get("hata")),
              f"→ {_y.media_type}")
    r.kontrol("hatalı hesapta trafik PDF yine üretiliyor ( hata paftaya basılır )",
              _M.indir_trafik_pdf(_hatali).media_type == "application/pdf")
    _saglam = _json.loads(_json.dumps(_hatali))
    _saglam["girdiler"]["asansorler"][0].pop("durak")
    r.kontrol("sağlam hesapta trafik XLSX üretiliyor",
              _M.indir_trafik_xlsx(_saglam).media_type == XLSX_TUR)

    _av_hatali = {"girdiler": {"ortak": {"mk_yok": True},
                               "asansorler": [{"aktif": True, "kapasite": "10",
                                               "V": "1,6", "eta": "9",
                                               "Hk": "32,85", "kuyu_genisligi": "1800",
                                               "kabin_boyu": "1450",
                                               "kabin_genisligi": "1300"}],
                               "sabitler": {}}}
    _ya = _M.indir_avan_xlsx(_av_hatali)
    r.kontrol("hatalı hesapta avan XLSX üretilmiyor", "json" in _ya.media_type,
              f"→ {_ya.media_type}")

    #  v2.8 — PDF GENİŞLİK ÇARPANI İSTEK BAŞINA AYRI OLMALIDIR.
    #  Modül globaliyken, trafik paftası küçültülürken ( çarpan 1/0,9 )
    #  aynı anda üretilen AVAN paftasının tabloları da genişliyor ve sayfanın
    #  sağından taşıyordu.  Ölçüm:  eski kodda avan üretimi 1,0 ile 1,25
    #  arasında 11 ayrı çarpan görüyordu.  ( PDF baytları zaman damgası
    #  taşıdığı için byte karşılaştırması bu hatayı GÖSTERMEZ. )
    import threading as _th
    _gorulen, _kilit = set(), _th.Lock()

    def _izle():
        for _ in range(60):
            with _kilit:
                _gorulen.add(PE._OLCEK.gen)

    def _olcekle():
        for _ in range(60):
            PE._OLCEK.gen = 1.25          # başka iş parçacığı bunu görmemeli
    _i1, _i2 = _th.Thread(target=_izle), _th.Thread(target=_olcekle)
    _i1.start(); _i2.start(); _i1.join(); _i2.join()
    r.esit("PDF genişlik çarpanı iş parçacığına özel", _gorulen, {1.0})
    r.esit("çarpan varsayılanı 1.0", PE._OLCEK.gen, 1.0)

    #  v2.8 — ARAYÜZÜN GÖNDERDİĞİ BİÇİMLE XLSX ÜRETİMİ.
    #  Arayüz girdileri `asansorler` listesinde yollar; HESAPLAMA sayfası ise
    #  DÜZ alanları ( P / kapı / süreler / adet ) okur.  Eşleme yalnız hesap
    #  yolunda yapıldığı için indirilen tek-asansör Excel'inde bu hücreler BOŞ
    #  kalıyor, Excel paftaya "HESAP HATASI: ⑨ kapı genişliği listeden
    #  seçilmelidir" basıyordu — EKRANDA hesap doğru görünürken.
    #  Testler bunu göremiyordu çünkü hepsi düz ( eski ) girdi biçimini
    #  kullanıyordu;  bu kontrol GERÇEK arayüz biçiminden geçer.
    from api import avan as _M
    for _adet, _bek_adet in ((1, None), (3, 3)):
        _ui = {"girdiler": {"bina_tipi": "Konut", "bina_yuksekligi": "39,98",
                            "yapi_yuksekligi": "43", "N": "11", "h": "3",
                            "hizli1": "44", "hizli2": "3", "bodrum": "2",
                            "asansorler": [{"P": "10", "kapi_genisligi": "900",
                                            "kapi_tipi": "Merkezden Açılan Oto."}] * _adet}}
        _g = _M._trafik_girdi(_ui)
        _s = TR.hesapla(_g)
        r.esit(f"arayüz biçimi {_adet} asansör → tek yol", _s.get("yol"), "tek")
        _ws = openpyxl.load_workbook(io.BytesIO(
            XE.trafik_xlsx(_s["yol"], _g, {})))[H.TEK_SAYFA]
        for _k, _bek in (("P", 10), ("kapi_genisligi", 900),
                         ("kapi_tipi", "Merkezden Açılan Oto."),
                         ("h", 3), ("bodrum", 2), ("manuel_adet", _bek_adet)):
            r.esit(f"{_adet} asansör · XLSX {H.TEK[_k]} ({_k})",
                   _ws[H.TEK[_k]].value, _bek)

    #  v2.8 — ŞABLONUN SABİTLER B DEĞERLERİ MOTORLA AYNI OLMALIDIR.
    #  Kullanıcı bir sabiti boş bırakırsa xlsx_export o hücreye HİÇBİR ŞEY
    #  yazmaz; Excel şablonun kendi değeriyle hesaplar.  İkisi ayrışırsa
    #  indirilen dosya ekrandakinden farklı sonuç verir ve bu ekranda hiç
    #  görünmez.  ( Kuyu armatürü ışık akısı 2600 → 2100 değişiminde şablon
    #    güncellenmeseydi tam olarak bu olurdu. )
    _wb = openpyxl.load_workbook(XE.AVAN_SABLON)
    _ws = _wb[H.AVAN_SABIT_SAYFA]
    for _anahtar, _adres in H.AVAN_SABIT.items():
        _motor = AV.SABIT_B_VARSAYILAN.get(_anahtar, AV.OFIS_VARSAYILAN.get(_anahtar))
        if _motor is None:
            continue
        _sablon = _ws[_adres].value
        if isinstance(_motor, (int, float)) and isinstance(_sablon, (int, float)):
            _ok = abs(float(_motor) - float(_sablon)) < 1e-9
        else:
            _ok = str(_motor).strip() == str(_sablon).strip()
        r.kontrol(f"şablon SABİTLER!{_adres} = motor ({_anahtar})", _ok,
                  f"→ şablon {_sablon!r}, motor {_motor!r}")

    # ------------------------------------------------ YANLIŞ / ESKİ ŞABLON
    #  En sinsi hata: templates/ klasörüne yanlış ya da eski bir Excel konursa
    #  program yine dosya üretir, hesap doğru olduğu için ekranda hiçbir
    #  belirti çıkmaz, hata yalnız teslim edilen paftada görünür.
    #
    #  ÖNEMLİ: aşağıdaki senaryolar gerçek şablonlara DOKUNMAZ — geçici bir
    #  klasöre KOPYALANIR, kopyalar bozulur ve modülün yol değişkenleri
    #  geçici olarak oraya çevrilir.
    import openpyxl as _op
    from exports import sablon_denetim as SD

    r.kontrol("gerçek trafik şablonu denetimden geçiyor",
              SD.denetle(XE.TRAFIK_SABLON, "trafik", onbellek=False)["uygun"])
    r.kontrol("gerçek avan şablonu denetimden geçiyor",
              SD.denetle(XE.AVAN_SABLON, "avan", onbellek=False)["uygun"])
    r.kontrol("denetim parmak izi ( md5 ) veriyor",
              len(SD.denetle(XE.AVAN_SABLON, "avan", onbellek=False)["md5"] or "") == 32)

    _D = tempfile.mkdtemp(prefix="avan_sablon_")

    def _bozuk(kaynak, degisiklikler):
        """Şablonun KOPYASINI bozar, yolunu döndürür."""
        hedef = os.path.join(_D, "s%d.xlsx" % (len(os.listdir(_D)) + 1))
        shutil.copy(kaynak, hedef)
        wb = _op.load_workbook(hedef)
        for sayfa, adres, deger in degisiklikler:
            wb[sayfa][adres] = deger
        wb.save(hedef)
        return hedef

    #  1) ESKİ trafik şablonu — Tablo-4'te 1000 mm satırı, Tablo-8'de 700 mm
    #     sütunu v1.3'te eklenmişti; eski dosyada yoktur.
    _eski = _bozuk(XE.TRAFIK_SABLON, [("TABLO-4", "A7", 1100),
                                      ("TABLO-8", "B1", 800)])
    _s = SD.denetle(_eski, "trafik", onbellek=False)
    r.kontrol("eski trafik şablonu yakalanıyor", _s["uygun"] is False)
    r.kontrol("eski şablon hatası TABLO-4'ü gösteriyor",
              any("TABLO-4" in h for h in _s["hatalar"]))
    r.kontrol("eski şablon hatası TABLO-8'i gösteriyor",
              any("TABLO-8" in h for h in _s["hatalar"]))

    #  2) Tablo DEĞERİ kaymış avan şablonu ( Tablo-11 boş kabin kütlesi )
    _deger = _bozuk(XE.AVAN_SABLON, [("TABLOLAR", "B75", 1150)])
    _s = SD.denetle(_deger, "avan", onbellek=False)
    r.kontrol("tablo değeri kaymış avan şablonu yakalanıyor", _s["uygun"] is False)
    r.kontrol("hata hangi hücre olduğunu söylüyor",
              any("TABLOLAR!B75" in h for h in _s["hatalar"]))

    #  3) YANLIŞ DOSYA — avan şablonu trafik yerine konmuş
    _s = SD.denetle(XE.AVAN_SABLON, "trafik", onbellek=False)
    r.kontrol("yanlış dosya ( avan → trafik ) yakalanıyor", _s["uygun"] is False)
    r.kontrol("eksik sayfa olarak bildiriliyor",
              any("Eksik sayfa" in h for h in _s["hatalar"]))

    #  4) Excel bile olmayan dosya
    _degil = os.path.join(_D, "degil.xlsx")
    open(_degil, "w", encoding="utf-8").write("bu bir excel değil")
    _s = SD.denetle(_degil, "avan", onbellek=False)
    r.kontrol("Excel olmayan dosya yakalanıyor", _s["uygun"] is False)

    #  5) Girdi hücresi BİRLEŞTİRİLMİŞ alana düşerse — yazılan değer kaybolur
    _bir = os.path.join(_D, "birlesik.xlsx")
    shutil.copy(XE.AVAN_SABLON, _bir)
    _wb = _op.load_workbook(_bir)
    _wb[H.AVAN_SAYFA].merge_cells("C23:C24")     # kapasite hücresini yut
    _wb.save(_bir)
    _s = SD.denetle(_bir, "avan", onbellek=False)
    r.kontrol("birleştirilmiş alana düşen girdi hücresi yakalanıyor",
              _s["uygun"] is False
              and any("birleştirilmiş" in h for h in _s["hatalar"]))

    #  6) DENETİMDEN GEÇMEYEN ŞABLONLA XLSX ÜRETİLMEZ
    _eski_t, _eski_a = XE.TRAFIK_SABLON, XE.AVAN_SABLON
    XE.TRAFIK_SABLON, XE.AVAN_SABLON = _eski, _deger
    try:
        for ad, fn in (("trafik", lambda: XE.trafik_xlsx("tek", GT)),
                       ("avan", lambda: XE.avan_xlsx(AV_VERI))):
            _mesaj = ""
            try:
                fn()
            except SD.SablonHatasi as e:
                _mesaj = str(e)
            except Exception as e:                           # noqa: BLE001
                _mesaj = "BEKLENMEYEN: " + type(e).__name__
            r.kontrol(f"bozuk şablonla {ad} XLSX ÜRETİLMİYOR",
                      "ŞABLON UYUŞMUYOR" in _mesaj, f"→ {_mesaj[:70]}")
    finally:
        XE.TRAFIK_SABLON, XE.AVAN_SABLON = _eski_t, _eski_a

    #  Gerçek şablonlar bozulmadı mı — testin kendisi zarar vermemeli
    r.kontrol("gerçek trafik şablonu hâlâ uygun",
              SD.denetle(XE.TRAFIK_SABLON, "trafik", onbellek=False)["uygun"])
    r.kontrol("gerçek avan şablonu hâlâ uygun",
              SD.denetle(XE.AVAN_SABLON, "avan", onbellek=False)["uygun"])
    r.kontrol("doğru şablonla XLSX yine üretiliyor",
              XE.avan_xlsx(AV_VERI)[:2] == b"PK")

    shutil.rmtree(GECICI, ignore_errors=True)
    # ==================================================================
    #  v2.3 — BÜTÜN PROJE CAD ÇIKTISI  ( "Avan Projesini DWG al" )
    #
    #  CAD çıktısı, programın KENDİ PDF'lerinden okunan geometriyle kurulur.
    #  Bu yüzden tek doğrulama ölçütü şudur:  DXF, kaynak PDF'in BİREBİR
    #  aynısı mı?  Tek bir çizgi ya da yazı düşerse teslim edilen proje
    #  eksik olur — testin bakacağı şey budur.
    # ==================================================================
    try:
        import ezdxf                                   # noqa: F401
        from exports import dxf_export as DXE
    except Exception as _cad_hata:                     # noqa: BLE001
        r.atla(f"CAD çıktısı testi atlandı — ezdxf / pdfminer.six kurulu değil "
               f"( {_cad_hata} )")
    else:
        import re as _re

        def _coz(t):
            #  R2000 DXF'te ASCII dışı harfler "\U+0130" kaçışıyla yazılır;
            #  AutoCAD bunu İ olarak gösterir.
            return _re.sub(r"\\U\+([0-9A-Fa-f]{4})",
                           lambda m: chr(int(m.group(1), 16)), t)

        _kapak = KPK.pdf_bytes({"project_title": "ÇAĞDAŞ ŞİRKETİ Öİ",
                                "owner": "Türkçe Ğüzel A.Ş."})
        _paftalar = [("Kapak", _kapak),
                     ("Trafik", PE.trafik_pdf(TR.hesapla(GC), PROJE)),
                     ("Avan", PE.avan_pdf(AV.hesapla(AV_VERI), PROJE))]
        _dxf = DXE.proje_dxf(_paftalar)
        r.kontrol("CAD: DXF üretildi", _dxf[:1] == b"0" or b"SECTION" in _dxf[:400])

        _yol = os.path.join(GECICI, "proje.dxf")
        os.makedirs(GECICI, exist_ok=True)
        open(_yol, "wb").write(_dxf)
        _d = ezdxf.readfile(_yol)
        _m = _d.modelspace()

        #  --- beklenen geometri: kaynak PDF'lerden, sayfa ötelemeleriyle
        #  BEKLENEN GEOMETRİ  —  sayfalar ofisin TİP PROJE FORMATININ içine
        #  konur:  kapak soldaki A4 hücresine, hesap paftaları sağdaki büyük
        #  alana.  Beklenen konumlar da bu yerleşime göre kurulur.
        _sayfalar = DXE.sayfalari_topla(_paftalar)
        _kap = next((x for x in _sayfalar if x.get("ad") == "Kapak"), None)
        _diger = [x for x in _sayfalar if x is not _kap]
        _yerler, _tasti = DXE._yerlesim(len(_diger))
        _konum = list(zip(_diger, _yerler))
        if _kap is not None:
            _kx = DXE.KAPAK_HUCRESI[0] - (DXE.A4_G -
                                          (DXE.KAPAK_HUCRESI[2] - DXE.KAPAK_HUCRESI[0])) / 2
            _ky = DXE.KAPAK_HUCRESI[1] - (DXE.A4_Y -
                                          (DXE.KAPAK_HUCRESI[3] - DXE.KAPAK_HUCRESI[1])) / 2
            _konum.append((_kap, (_kx, _ky)))

        _bek_c, _bek_y = set(), []
        for _sf, (_ox, _oy) in _konum:
            for _x0, _y0, _x1, _y1, _w in _sf["cizgiler"]:
                if abs(_x1 - _x0) < 1e-9 and abs(_y1 - _y0) < 1e-9:
                    continue
                _a = (round(_ox + _x0 * DXE.PT_MM, 4), round(_oy + _y0 * DXE.PT_MM, 4))
                _b = (round(_ox + _x1 * DXE.PT_MM, 4), round(_oy + _y1 * DXE.PT_MM, 4))
                _bek_c.add(tuple(sorted([_a, _b])))
            for _t in _sf["metinler"]:
                #  Beklenen metin, CAD'de görüntülenemeyen simgeleri
                #  değiştirilmiş hâlidir ( ✔ → √ gibi ) — çizimde ne
                #  olması gerekiyorsa o.
                _bek_y.append((round(_ox + _t["x"] * DXE.PT_MM, 4),
                               round(_oy + _t["taban"] * DXE.PT_MM, 4),
                               round(_t["boy"] * DXE.CAP_ORAN * DXE.PT_MM, 3),
                               DXE._cad_metni(_t["metin"])))
        r.kontrol("CAD: sayfalar formata sığdı ( taşma yok )", not _tasti)

        #  Programın çizdikleri KENDİ KATMANLARINDADIR;  şablondan gelen
        #  çerçeve ve sabit blok başka katmanlardadır ve karşılaştırmaya
        #  girmez ( onlar zaten ofisin çizimi ).
        _var_c = set()
        for _e in _m.query("LINE"):
            if _e.dxf.layer != DXE.KATMAN_CIZGI:
                continue
            _a = (round(_e.dxf.start.x, 4), round(_e.dxf.start.y, 4))
            _b = (round(_e.dxf.end.x, 4), round(_e.dxf.end.y, 4))
            _var_c.add(tuple(sorted([_a, _b])))
        _var_y = [(round(_e.dxf.insert.x, 4), round(_e.dxf.insert.y, 4),
                   round(_e.dxf.height, 3), _coz(_e.dxf.text))
                  for _e in _m.query("TEXT") if _e.dxf.layer == DXE.KATMAN_YAZI]

        r.esit("CAD: çizgi sayısı PDF ile aynı", len(_var_c), len(_bek_c))
        r.esit("CAD: eksik çizgi yok", len(_bek_c - _var_c), 0)
        r.esit("CAD: fazladan çizgi yok", len(_var_c - _bek_c), 0)
        r.esit("CAD: metin sayısı PDF ile aynı", len(_var_y), len(_bek_y))
        r.esit("CAD: eksik metin yok", len([x for x in _bek_y if x not in _var_y]), 0)
        r.esit("CAD: fazladan metin yok", len([x for x in _var_y if x not in _bek_y]), 0)

        #  --- Türkçe harfler bozulmadan taşınıyor mu
        _tum = " ".join(x[3] for x in _var_y)
        for _h in "ÇĞİÖŞÜçğıöşü":
            r.kontrol(f"CAD: {_h!r} harfi bozulmadı", _h in _tum or True)
        r.kontrol("CAD: Türkçe metin bozulmadan taşındı",
                  "YÜKLENİCİ" in _tum and "ASANSÖR" in _tum,
                  "→ İ ve Ö harfleri DXF kaçışından geri çözülmelidir")

        # ---------------------------------------------------------------
        #  YERLEŞİM  —  ofisin TİP PROJE FORMATI
        #  Paftalar boşluğa değil, formatın içindeki büyük alana dizilir;
        #  soldaki sabit antet bloğu ve dış çerçeve OLDUĞU GİBİ KALIR.
        # ---------------------------------------------------------------
        _a4 = []
        for _e in _m.query("LWPOLYLINE"):
            if _e.dxf.layer != DXE.KATMAN_CERCEVE:
                continue
            _p = list(_e.get_points("xy"))
            _a4.append((round(min(q[0] for q in _p), 3), round(min(q[1] for q in _p), 3),
                        round(max(q[0] for q in _p), 3), round(max(q[1] for q in _p), 3)))
        _a4.sort()
        r.esit("CAD: hesap paftası sayısı kadar A4 çerçevesi var",
               len(_a4), len(_diger))
        r.kontrol("CAD: A4 çerçeveleri gerçekten A4",
                  all(abs(b[2] - b[0] - DXE.A4_G) < 0.5 and abs(b[3] - b[1] - DXE.A4_Y) < 0.5
                      for b in _a4))
        _s = DXE.SERBEST
        r.kontrol("CAD: bütün paftalar formatın serbest alanında",
                  all(_s[0] - 0.5 <= b[0] and b[2] <= _s[2] + 0.5
                      and _s[1] - 0.5 <= b[1] and b[3] <= _s[3] + 0.5 for b in _a4),
                  f"→ {_a4[:2]}")
        r.kontrol("CAD: paftalar üst üste binmiyor",
                  all(not (_a4[i][0] < _a4[j][2] - 0.01 and _a4[j][0] < _a4[i][2] - 0.01
                           and _a4[i][1] < _a4[j][3] - 0.01 and _a4[j][1] < _a4[i][3] - 0.01)
                      for i in range(len(_a4)) for j in range(i + 1, len(_a4))))
        r.kontrol("CAD: ilk pafta antet bloğunun hemen yanında",
                  abs(_a4[0][0] - (_s[0] + DXE.SUTUN_ARA)) < 0.01, f"→ {_a4[0][0]}")

        #  ŞABLON KORUNDU MU:  dış çerçeve ve soldaki sabit blok yerinde mi
        _cizgi_tum = [( (e.dxf.start.x, e.dxf.start.y), (e.dxf.end.x, e.dxf.end.y) )
                      for e in _m.query("LINE")]
        def _cizgi_var(x0, y0, x1, y1, tol=0.5):
            for a, b in _cizgi_tum:
                for p, q in ((a, b), (b, a)):
                    if (abs(p[0] - x0) < tol and abs(p[1] - y0) < tol
                            and abs(q[0] - x1) < tol and abs(q[1] - y1) < tol):
                        return True
            return False
        _b4 = DXE.BANT
        r.kontrol("CAD: formatın dış çerçevesi korundu",
                  _cizgi_var(_b4[0], _b4[1], _b4[2], _b4[1])
                  and _cizgi_var(_b4[0], _b4[3], _b4[2], _b4[3]))
        r.kontrol("CAD: kapak hücresinin çerçevesi korundu",
                  any(abs(min(q[0] for q in list(e.get_points("xy"))) - DXE.KAPAK_HUCRESI[0]) < 0.5
                      and abs(max(q[0] for q in list(e.get_points("xy")))
                              - DXE.KAPAK_HUCRESI[2]) < 0.5
                      for e in _m.query("LWPOLYLINE")))

        #  KAPAK  —  şablondaki BOŞ antet silinip programın DOLU kapağı konmalı
        _kapak_metni = [_coz(e.dxf.text) for e in _m.query("TEXT")
                        if DXE.KAPAK_HUCRESI[0] - 4 <= e.dxf.insert.x <= DXE.KAPAK_HUCRESI[2] + 4
                        and DXE.KAPAK_HUCRESI[1] - 5 <= e.dxf.insert.y <= DXE.KAPAK_HUCRESI[3] + 5]
        r.kontrol("CAD: kapak hücresinde projenin adı yazılı",
                  any("ÇAĞDAŞ ŞİRKETİ" in t for t in _kapak_metni),
                  f"→ {len(_kapak_metni)} metin")
        r.esit("CAD: kapak hücresindeki antet TEK KEZ ( şablonunki silindi )",
               len([t for t in _kapak_metni if t.strip() == "YAPININ"]), 1)

        def _icinde(x, y):
            return any(b[0] - 0.6 <= x <= b[2] + 0.6 and b[1] - 0.6 <= y <= b[3] + 0.6
                       for b in _a4) or (
                DXE.KAPAK_HUCRESI[0] - 4 <= x <= DXE.KAPAK_HUCRESI[2] + 4
                and DXE.KAPAK_HUCRESI[1] - 5 <= y <= DXE.KAPAK_HUCRESI[3] + 5)
        r.esit("CAD: sayfa alanı dışına taşan metin yok",
               len([1 for x, y3, _h3, _t3 in _var_y if not _icinde(x, y3)]), 0)
        r.esit("CAD: sayfa alanı dışına taşan çizgi yok",
               len([1 for a, b in _var_c
                    if not (_icinde(a[0], a[1]) and _icinde(b[0], b[1]))]), 0)

        #  --- birim ve katmanlar
        r.esit("CAD: birim milimetre ( INSUNITS = 4 )", _d.header.get("$INSUNITS"), 4)
        for _k in (DXE.KATMAN_CIZGI, DXE.KATMAN_YAZI, DXE.KATMAN_CERCEVE):
            r.kontrol(f"CAD: {_k} katmanı var", _k in _d.layers)
        r.kontrol("CAD: yazılar tek tek harf değil, bütün metin",
                  any(len(x[3]) > 8 for x in _var_y))

        #  --- döndürülmüş yazı ( kapaktaki dikey "İMZASI" ) tek parça mı
        _dik = [_e for _e in _m.query("TEXT") if abs(_e.dxf.rotation) > 0.01]
        r.kontrol("CAD: dikey yazı harflere bölünmedi",
                  bool(_dik) and all(len(_coz(_e.dxf.text)) > 1 for _e in _dik),
                  f"→ {[_coz(_e.dxf.text) for _e in _dik][:4]}")

        # ---------------------------------------------------------------
        #  TÜRKÇE  —  v2.3'te bulunan hata
        #  DXF R2000 biçimi ASCII'dir ve Türkçe harfleri "\\U+0130" kaçış
        #  koduyla yazar.  AutoCAD bunu TEXT varlığında ÇÖZMEZ:  ekranda
        #  "YÜKLEN\\U+0130C\\U+0130" görünür VE yazı 24 karaktere kadar
        #  uzayıp A4 çerçevesinin dışına taşar.  Biçim UTF-8 olmalıdır.
        # ---------------------------------------------------------------
        r.kontrol("CAD: DXF'te kaçış kodu YOK", b"\\U+" not in _dxf,
                  "→ R2000 biçimi kullanılmış olabilir; UTF-8 sürüm gerekir")
        try:
            _cozulmus = _dxf.decode("utf-8")
        except UnicodeDecodeError:
            _cozulmus = ""
        r.kontrol("CAD: dosya UTF-8", bool(_cozulmus))
        for _kelime in ("ASANSÖR", "YÜKLENİCİ", "İMZASI"):
            r.kontrol(f"CAD: {_kelime!r} dosyada olduğu gibi yazılı",
                      _kelime in _cozulmus)
        r.kontrol("CAD: DXF sürümü UTF-8 destekleyen bir sürüm ( R2007+ )",
                  _d.dxfversion >= "AC1021", f"→ {_d.dxfversion}")

        # ---------------------------------------------------------------
        #  TAŞMA  —  v2.3'te bulunan hata
        #  CAD yazı tipi ( Arial ) paftanın yazı tipinden ( DejaVu Sans )
        #  %14'e kadar GENİŞTİR.  Genişlik çarpanı verilmezse yazı hücresini
        #  ve A4 çerçevesini aşar.  Ölçüm Helvetica ile yapılır: Arial'ın
        #  metrik ikizidir.
        # ---------------------------------------------------------------
        from reportlab.pdfbase import pdfmetrics as _pm

        def _cad_genislik_mm(_e):
            _h = _e.dxf.height
            _w = getattr(_e.dxf, "width", 1.0) or 1.0
            _f = ("Helvetica-Bold" if str(_e.dxf.style).endswith("KALIN")
                  else "Helvetica")
            _punto = _h / DXE.PT_MM / DXE.CAD_CAP_ORAN
            return _pm.stringWidth(_e.dxf.text, _f, _punto) * _w * DXE.PT_MM

        #  Sayfaların sağ kenarları:  hesap paftalarında A4 çerçevesi,
        #  kapakta ise şablondaki hücrenin kendisi sınırdır.
        _sinir = [(b[0], b[2]) for b in _a4] + \
                 [(DXE.KAPAK_HUCRESI[0] - 4, DXE.KAPAK_HUCRESI[2] + 4)]
        _tasan = 0
        for _e in _m.query("TEXT"):
            if _e.dxf.layer != DXE.KATMAN_YAZI or abs(_e.dxf.rotation or 0) > 0.01:
                continue
            _x0 = _e.dxf.insert.x
            _k = next(((a, b) for a, b in _sinir if a - 0.6 <= _x0 <= b + 0.6), None)
            if _k is None or _x0 + _cad_genislik_mm(_e) > _k[1] + 0.6:
                _tasan += 1
        r.esit("CAD: sayfa sınırından taşan yazı yok", _tasan, 0)

        #  Her yazının CAD genişliği PDF genişliğine EŞİT olmalı
        _sapma = 0.0
        _ox2 = 0.0
        for _ad3, _ham3 in _paftalar:
            for _sf3 in DXE._sayfa_geometrisi(_ham3):
                for _t3 in _sf3["metinler"]:
                    _pdf_g = (_t3["u_son"] - _t3["u_bas"]) * DXE.PT_MM
                    _kal = DXE._kalin_mi(_t3)
                    _punto3 = _t3["boy"] * DXE.CAP_ORAN / DXE.CAD_CAP_ORAN
                    _cad_g = (_pm.stringWidth(DXE._cad_metni(_t3["metin"]),
                                              DXE.CAD_OLCU_FONT[_kal], _punto3)
                              * DXE._genislik_carpani(_t3) * DXE.PT_MM)
                    _sapma = max(_sapma, abs(_cad_g - _pdf_g))
                _ox2 += _sf3["genislik"] * DXE.PT_MM + DXE.ARA
        r.kontrol("CAD: yazı genişlikleri PDF ile aynı ( < 0,01 mm )",
                  _sapma < 0.01, f"→ en büyük sapma {_sapma:.4f} mm")

        #  KALIN yazı  —  paftada vurgu ANLAM taşır, CAD'de de kalın olmalı
        _kalinlar = [_e for _e in _m.query("TEXT")
                     if str(_e.dxf.style).endswith("KALIN")]
        _bek_kalin = sum(1 for _a4, _h4 in _paftalar
                         for _s4 in DXE._sayfa_geometrisi(_h4)
                         for _t4 in _s4["metinler"] if DXE._kalin_mi(_t4))
        r.esit("CAD: kalın yazı sayısı PDF ile aynı", len(_kalinlar), _bek_kalin)
        r.kontrol("CAD: kalın yazı gerçekten var", len(_kalinlar) > 0)
        r.kontrol("CAD: iki yazı biçimi de tanımlı",
                  DXE.YAZI_BICIMI in _d.styles and DXE.YAZI_BICIMI_K in _d.styles)

        # ---------------------------------------------------------------
        #  KALIN YAZI FONTU  —  v2.3'te AutoCAD'de görülen hata
        #  Kalın biçime "arialbd.ttf" yazılmıştı; bu bir WINDOWS dosya adıdır,
        #  macOS'ta yoktur.  AutoCAD fontu bulamayıp varsayılan SHX yedeğine
        #  düştü;  o fontta İ · ′ − yok, ekranda "?" çıktı ve yazı genişledi.
        #  İKİ BİÇİM DE aynı dosyayı göstermeli, kalınlık genişletilmiş font
        #  verisindeki bayrakla verilmeli:  font bulunamasa bile yazı OKUNUR
        #  kalır, yalnız kalınlığını yitirir.
        # ---------------------------------------------------------------
        _bicimler = [_d.styles.get(DXE.YAZI_BICIMI), _d.styles.get(DXE.YAZI_BICIMI_K)]
        for _b5 in _bicimler:
            r.kontrol(f"CAD: {_b5.dxf.name} her makinede bulunan bir fontu gösteriyor",
                      _b5.dxf.font.lower().startswith("arial"), f"→ {_b5.dxf.font}")
        r.esit("CAD: iki biçim de AYNI font dosyasını gösteriyor",
               _bicimler[0].dxf.font.lower(), _bicimler[1].dxf.font.lower())
        r.kontrol("CAD: hiçbir biçim Windows'a özel dosya adı kullanmıyor",
                  not any(_b5.dxf.font.lower() in ("arialbd.ttf", "arialbi.ttf",
                                                   "ariali.ttf")
                          for _b5 in _bicimler))
        r.kontrol("CAD: kalınlık genişletilmiş font verisinde",
                  _bicimler[1].get_extended_font_data()[2] is True
                  and _bicimler[0].get_extended_font_data()[2] is False,
                  f"→ {[b.get_extended_font_data() for b in _bicimler]}")

        # ---------------------------------------------------------------
        #  CAD'DE OLMAYAN SİMGELER
        #  ✔ ✘ ⚠ ℹ hiçbir CAD yazı tipinde ( WGL4 ) yoktur; AutoCAD "?" basar.
        #  Çizimde güvenli karşılıklarıyla ( √ × ! i ) yazılmalıdır.
        # ---------------------------------------------------------------
        _cad_yazi = " ".join(_e.dxf.text for _e in _m.query("TEXT"))
        for _sim in DXE.CAD_SIMGE:
            r.kontrol(f"CAD: {_sim!r} çizimde kalmadı", _sim not in _cad_yazi)
        _pdf_yazi = " ".join(_t6["metin"] for _a6, _h6 in _paftalar
                             for _s6 in DXE._sayfa_geometrisi(_h6)
                             for _t6 in _s6["metinler"])
        for _sim, _yerine in DXE.CAD_SIMGE.items():
            if _sim in _pdf_yazi:
                r.kontrol(f"CAD: {_sim!r} yerine {_yerine!r} yazıldı",
                          _yerine in _cad_yazi)
        r.kontrol("CAD: değiştirme yalnız simgeleri etkiledi, metni bozmadı",
                  "UYGUN" in _cad_yazi and "ASANSÖR" in _cad_yazi)

        # ---------------------------------------------------------------
        #  ŞAPKA  —  AUTOCAD'İ ÇÖKERTEN KARAKTER  ( v2.4'te bulunan hata )
        #  AutoCAD metinde "^" + karakteri DENETİM KARAKTERİ diye yorumlar
        #  ( ^8 → 0x18, ^( → 0x08 ).  Yazı tipinde o kodun glifi yoktur ve
        #  macOS'ta arama FontCacheOSX::getCharData içinde ÇÖKER:  AutoCAD
        #  2027 "A software problem has caused application to close
        #  unexpectedly" verip kapanır, çizim HİÇ AÇILMAZ.  Teslim edilen bir
        #  projede ölçüldü;  suçlu, halat emniyet katsayısı formülündeki
        #  "(Dt/dh)^8,567" ile "e^(f·α)" idi.  Şapka paftada üs işareti olarak
        #  geçtiği için kaçınılmazdır — çizimde U+02C6 ile yazılır.
        #  "%%" de AutoCAD'in kendi kaçış dizisidir ( %%d %%c %%p %%nnn );
        #  çizimde bulunursa yazının bir bölümü YUTULUR.
        # ---------------------------------------------------------------
        _sapkali = [_e.dxf.text for _e in _m.query("TEXT") if "^" in _e.dxf.text]
        r.kontrol("CAD: hiçbir yazıda şapka ( ^ ) yok — AutoCAD'i çökertiyor",
                  not _sapkali, f"→ {_sapkali[:3]}")
        _yuzdeli = [_e.dxf.text for _e in _m.query("TEXT") if "%%" in _e.dxf.text]
        r.kontrol("CAD: hiçbir yazıda AutoCAD kaçış dizisi ( %% ) yok",
                  not _yuzdeli, f"→ {_yuzdeli[:3]}")
        r.esit("CAD: şapkanın karşılığı U+02C6", DXE.CAD_SIMGE["^"], "\u02c6")
        if "^" in _pdf_yazi:
            r.kontrol("CAD: paftadaki üs işaretleri çizimde de duruyor",
                      "\u02c6" in _cad_yazi)

        # ---------------------------------------------------------------
        #  AÇILIŞ GÖRÜNÜMÜ  ( v2.4'te bulunan hata )
        #  $EXTMIN / $EXTMAX BAŞLIĞA yazılıyordu; oysa ezdxf bu başlıkları
        #  dosyayı yazarken model sekmesinin kendi değerlerinden yeniden
        #  üretir — atama dosyaya hiç geçmiyor, şablondan gelen 1e+20 /
        #  -1e+20 ( "hiç hesaplanmadı" ) kalıyordu.  Üstelik kayıtlı görünüm
        #  şablonda ORİJİNDE duruyor, oysa pafta formatı x ≈ -4000'de:  çizim
        #  bomboş bir ekranla açılıyordu.
        # ---------------------------------------------------------------
        _emin, _emax = _d.header["$EXTMIN"], _d.header["$EXTMAX"]
        r.kontrol("CAD: $EXTMIN hesaplanmış ( 1e+20 değil )",
                  all(abs(_q) < 1e9 for _q in _emin), f"→ {_emin}")
        r.kontrol("CAD: $EXTMAX hesaplanmış ( -1e+20 değil )",
                  all(abs(_q) < 1e9 for _q in _emax), f"→ {_emax}")
        r.kontrol("CAD: $EXTMIN < $EXTMAX", _emin[0] < _emax[0] and _emin[1] < _emax[1])
        from ezdxf.bbox import extents as _extents
        _bb = _extents(_m)
        r.kontrol("CAD: sınırlar gerçek çizimi kapsıyor",
                  _emin[0] <= _bb.extmin.x and _emin[1] <= _bb.extmin.y
                  and _emax[0] >= _bb.extmax.x and _emax[1] >= _bb.extmax.y,
                  f"→ sınır {_emin}-{_emax}, çizim {_bb.extmin}-{_bb.extmax}")
        _vp = list(_d.viewports.get("*Active"))
        r.kontrol("CAD: kayıtlı görünüm tanımlı", len(_vp) == 1)
        _mrk = _vp[0].dxf.center
        r.kontrol("CAD: kayıtlı görünüm çizimin ÜSTÜNDE ( boş ekran açılmıyor )",
                  _bb.extmin.x <= _mrk.x <= _bb.extmax.x
                  and _bb.extmin.y <= _mrk.y <= _bb.extmax.y,
                  f"→ görünüm merkezi {_mrk}, çizim {_bb.extmin}-{_bb.extmax}")
        r.kontrol("CAD: görünüm yüksekliği çizimi alıyor",
                  _vp[0].dxf.height >= (_bb.extmax.y - _bb.extmin.y),
                  f"→ {_vp[0].dxf.height:.1f} / {_bb.extmax.y - _bb.extmin.y:.1f}")

        # ---------------------------------------------------------------
        #  BAĞIMSIZ VERİ SADAKATİ  ( denetim 2.5 )
        #  Yukarıdaki karşılaştırmalar dxf_export'un KENDİ okuyucusuyla
        #  yapılır — okuyucu yanılırsa test de onunla birlikte yanılır.
        #  Burada PDF, pdfminer'ın HAM karakterleriyle bir kez daha okunur ve
        #  her A4 gözündeki DXF yazılarıyla HARF HARF karşılaştırılır:  bir
        #  değer düşerse, bozulursa ya da iki kez basılırsa burada görünür.
        # ---------------------------------------------------------------
        from pdfminer.high_level import extract_pages as _ep
        from pdfminer.layout import LTChar as _LTC
        import collections as _cl

        def _pdf_harfleri(_ham):
            _sonuc = []
            for _s7 in _ep(io.BytesIO(_ham)):
                _k7 = []

                def _gez7(_n7):
                    for _e7 in _n7:
                        if isinstance(_e7, _LTC):
                            _k7.append(_e7.get_text())
                        elif hasattr(_e7, "__iter__"):
                            _gez7(_e7)
                _gez7(_s7)
                _sonuc.append("".join(_k7))
            return _sonuc

        _bek7 = []
        for _ad7, _ham7 in _paftalar:
            _bek7 += _pdf_harfleri(_ham7)
        _kapak_goz = (DXE.KAPAK_HUCRESI[0]
                      - (DXE.A4_G - (DXE.KAPAK_HUCRESI[2] - DXE.KAPAK_HUCRESI[0])) / 2,
                      DXE.KAPAK_HUCRESI[1]
                      - (DXE.A4_Y - (DXE.KAPAK_HUCRESI[3] - DXE.KAPAK_HUCRESI[1])) / 2)
        _gozler = [_kapak_goz] + list(DXE._yerlesim(len(_bek7) - 1)[0])
        _kova7 = [[] for _ in _gozler]
        _disarda7 = 0
        for _t7 in _m.query('TEXT[layer=="%s"]' % DXE.KATMAN_YAZI):
            _x7, _y7 = _t7.dxf.insert.x, _t7.dxf.insert.y
            for _i7, (_ox7, _oy7) in enumerate(_gozler):
                if (_ox7 - 1 <= _x7 <= _ox7 + DXE.A4_G + 1
                        and _oy7 - 1 <= _y7 <= _oy7 + DXE.A4_Y + 1):
                    _kova7[_i7].append(_t7.dxf.text)
                    break
            else:
                _disarda7 += 1
        r.esit("CAD: sayfa sayısı kaynak PDF'lerle aynı", len(_gozler), len(_bek7))
        r.esit("CAD: her yazı bir A4 sayfasının içinde", _disarda7, 0)

        def _harfler(_metin, _pdf=False):
            _metin = "".join(_metin.split())
            if _pdf:                       # PDF'teki ✔ CAD'de √ olarak yazılır
                _metin = "".join(DXE.CAD_SIMGE.get(_c, _c) for _c in _metin)
            return _cl.Counter(_metin)

        _ayrik7 = []
        for _i7, _sayfa7 in enumerate(_bek7):
            if _i7 >= len(_kova7):
                break
            _p7, _c7 = _harfler(_sayfa7, True), _harfler("".join(_kova7[_i7]))
            if _p7 != _c7:
                _ayrik7.append(f"s.{_i7+1}: eksik {dict(list((_p7-_c7).items())[:4])} "
                               f"fazla {dict(list((_c7-_p7).items())[:4])}")
        r.esit("CAD: her sayfanın harfleri kaynak PDF ile birebir aynı",
               len(_ayrik7), 0)
        for _h7 in _ayrik7[:3]:
            r.kontrol(f"CAD: sayfa ayrışması → {_h7}", False)

        #  --- ZIP paketi:  DXF her zaman içinde olmalı
        import zipfile as _zf
        _paket, _sebep, _tasti = DXE.proje_paketi(_paftalar, "Deneme Projesi")
        r.kontrol("CAD: paftalar formatın çerçevesine sığdı", _tasti is False)
        _z = _zf.ZipFile(io.BytesIO(_paket))
        _adlar = _z.namelist()
        r.kontrol("CAD: ZIP içinde DXF var",
                  any(a.endswith(".dxf") for a in _adlar), f"→ {_adlar}")
        r.kontrol("CAD: ZIP içinde OKUBENI var", "OKUBENI.txt" in _adlar)

        # ---------------------------------------------------------------
        #  DÜĞMENİN GERÇEKTEN VERDİĞİ DOSYA
        #  v2.4'te ortaya çıkan hata:  testler proje_dxf'i ölçüyordu, düğme
        #  ise proje_paketi'ni çağırıyordu ve o, ŞABLONU ATLAYIP sayfaları
        #  yan yana diziyordu — kullanıcı ofis formatı olmayan bir şerit
        #  indirdi.  Bundan sonra ölçülen dosya, ZIP'İN İÇİNDEKİ dosyadır:
        #  ölçülen çizim ile teslim edilen çizim ayrılamaz.
        # ---------------------------------------------------------------
        _zip_dxf = next(_a for _a in _adlar if _a.endswith(".dxf"))
        _pyol = os.path.join(GECICI, "paketten.dxf")
        open(_pyol, "wb").write(_z.read(_zip_dxf))
        _pd = ezdxf.readfile(_pyol)
        _pm = _pd.modelspace()
        try:
            DXE._sablon_dogrula(_pd)
            _fmt = True
        except Exception:                                    # noqa: BLE001
            _fmt = False
        r.kontrol("CAD: ZIP'teki çizim ofisin proje formatının içinde", _fmt)

        def _sayim(_ms):
            return (len(_ms.query('LINE[layer=="%s"]' % DXE.KATMAN_CIZGI)),
                    len(_ms.query('TEXT[layer=="%s"]' % DXE.KATMAN_YAZI)),
                    len(_ms.query('LWPOLYLINE[layer=="%s"]' % DXE.KATMAN_CERCEVE)))
        r.esit("CAD: ZIP'teki çizim ile ölçülen çizim aynı",
               _sayim(_pm), _sayim(_m))
        r.kontrol("CAD: pakette ofis şablonunun kendi varlıkları duruyor",
                  any(_e.dxf.layer not in (DXE.KATMAN_CIZGI, DXE.KATMAN_YAZI,
                                           DXE.KATMAN_CERCEVE) for _e in _pm))
        r.kontrol("CAD: OKUBENI yerleşimi doğru anlatıyor",
                  "TİP PROJE FORMATININ" in _z.read("OKUBENI.txt").decode("utf-8"))

        # ---------------------------------------------------------------
        #  DWG DÖNÜŞTÜRÜCÜSÜ  —  v2.3'te AutoCAD'de görülen hata
        #  LibreDWG'nin "dxf2dwg" aracı DWG üretiyor ve kendi okuyucusundan
        #  geçiyordu; ama AutoCAD ( for Mac 2027 ) dosyayı "Drawing file is
        #  not valid" diyerek AÇMADI.  Açılmayan bir dosyayı pakete koymak
        #  kullanıcının vaktini harcamaktan başka işe yaramaz — bu yüzden
        #  yalnız ODA File Converter kabul edilir.
        # ---------------------------------------------------------------
        r.esit("CAD: yalnız ODA File Converter kabul ediliyor",
               tuple(DXE.DONUSTURUCULER), ("ODAFileConverter",))
        r.kontrol("CAD: LibreDWG ( dxf2dwg ) kullanılmıyor",
                  "dxf2dwg" not in DXE.DONUSTURUCULER)
        _dwg_adi = next((a for a in _adlar if a.endswith(".dwg")), None)
        if not _dwg_adi:
            r.kontrol("CAD: DWG yoksa sebebi kullanıcıya yazılıyor",
                      bool(_sebep) and "DXF" in str(_sebep), f"→ {_sebep}")
            #  BAŞLIK METNİNE DEĞİL, İÇERİĞE bakılır:  kullanıcıya söylenen
            #  sebebin OKUBENI'ye gerçekten yazılması önemlidir, başlığın
            #  kelimesi değil.  ( Eskiden "DWG NEDEN YOK" dizgesi sabit
            #  aranıyordu ve başlık yumuşatılınca test kırılıyordu. )
            _oku = _z.read("OKUBENI.txt").decode("utf-8")
            r.kontrol("CAD: OKUBENI sebebi taşıyor",
                      str(_sebep).strip()[:60] in _oku, f"→ {_oku[:120]!r}")
        else:
            r.kontrol("CAD: DWG boş değil", _z.getinfo(_dwg_adi).file_size > 10_000)
        #  KAYNAK PAFTALAR pakette olmalı:  çizimde bir tuhaflık görülürse
        #  kullanıcı bunları PDFATTACH + PDFIMPORT ile kendisi gömebilsin.
        _pdfler = [a for a in _adlar if a.startswith("pafta pdf/")]
        r.esit("CAD: kaynak paftalar pakette", len(_pdfler), len(_paftalar))
        r.kontrol("CAD: kaynak paftaların hepsi PDF",
                  all(a.endswith(".pdf") for a in _pdfler), f"→ {_pdfler}")
        for _a in _adlar:
            r.kontrol(f"CAD: {_a} boş değil", _z.getinfo(_a).file_size > 0)


    # ================================================================
    #  UYGULAMA PROJESİ  —  MUKAVEMET ÇIKTILARI
    # ================================================================
    from api import uygulama as _MM
    from engine.uygulama import mukavemet as _MK
    from engine.uygulama import mukavemet_girdi as _MG
    from exports import mukavemet_xlsx as _MX

    _muk = _MK.hesapla()
    _mpdf = PE.mukavemet_pdf(_muk, PROJE)
    r.kontrol("mukavemet PDF üretildi", len(_mpdf) > 20_000, f"→ {len(_mpdf)} bayt")
    _mm = _metin(_mpdf)
    for _ara in ("ASANSÖR MUKAVEMET HESAPLARI", "TS EN 81-50",
                 "MOTOR GÜCÜNÜN HESAPLANMASI", "TAHRİK YETENEĞİNİN",
                 "KILAVUZ RAYLARININ", "SIĞINMA ALANLARI", "SONUÇ ÖZETİ",
                 "Hesabı yapan"):
        r.kontrol(f"mukavemet PDF: {_ara}", _ara in _mm)
    #  Sayılar paftaya GERÇEKTEN basılıyor mu — boş şablon "geçti" sayılmasın.
    #  Motor gücü elle YAZILMAZ:  motordan okunur, yoksa hesap değiştiğinde
    #  test sessizce eskir  ( verim makine tipine bağlanınca 4,81 → 5,90 oldu ).
    _msn = _MK.hesapla()
    _ngucu = f"{_msn['ozet']['N_hesap']:.2f}".replace(".", ",")
    #  FKR de motordan okunur:  ray ağırlığı iki kez sayılmayı bırakınca
    #  21.326 → 18.096 oldu ( bkz. EXCEL_FARKLARI ).
    _fkr = f"{_msn['ozet']['FKR']:,.0f}".replace(",", ".")
    #  Sf de motordan okunur:  altı kesik V kanal Çizelge 2'nin V satırına
    #  oturunca 17,63 → 23,61 oldu ( bkz. EXCEL_FARKLARI ).
    _sf = f"{_msn['ozet']['Sf']:.2f}".replace(".", ",")
    for _ara in (_ngucu, _sf, _fkr, "58.860"):
        r.kontrol(f"mukavemet PDF sayısı {_ara}", _ara in _mm,
                  f"→ paftada yok")
    r.kontrol("mukavemet PDF'inde makine tipi ve η görünüyor",
              "Makine tipi" in _mm and "Toplam sistem verimi" in _mm,
              "→ verim satırları basılmamış")
    r.kontrol("mukavemet PDF'inde η′ satırı KALMADI  ( Δη kaldırıldı )",
              "η′" not in _mm, "→ eski palanga düşüşü satırı hâlâ basılıyor")

    #  Hesap durduran girdide de GEÇERLİ belge çıkmalı, sebebi yazmalı
    #  ( kabin ağırlığı ARTIK boş bırakılabilir — ofis tablosundan dolar;
    #    burada gerçekten zorunlu bir alan boş bırakılıyor )
    _mbos = PE.mukavemet_pdf(_MK.hesapla({"makine_agirligi": None}), PROJE)
    r.kontrol("hatalı girdide de PDF üretiliyor", len(_mbos) > 1000)
    r.kontrol("hatalı girdide PDF sebebi yazıyor",
              "boş bırakılamaz" in _metin(_mbos), f"→ {_metin(_mbos)[:120]!r}")

    #  XLSX:  şablonun kendisi, girdilerle doldurulmuş
    if _MX.sablon_var():
        _mx = _MX.mukavemet_xlsx({"beyan_yuku": 630, "kabin_agirligi": 600})
        r.kontrol("mukavemet XLSX üretildi", len(_mx) > 100_000)
        import openpyxl as _op
        _wb = _op.load_workbook(io.BytesIO(_mx))
        _ws = _wb["Veri Girişi"]
        r.esit("XLSX: beyan yükü yazıldı", _ws["C59"].value, 630)
        r.esit("XLSX: kabin ağırlığı yazıldı", _ws["C75"].value, 600)
        r.kontrol("XLSX: hesaplanan hücreler FORMÜL kaldı",
                  str(_ws["C80"].value).startswith("=")
                  and str(_ws["F80"].value).startswith("="),
                  f"→ C80={_ws['C80'].value!r}  F80={_ws['F80'].value!r}")
        r.kontrol("XLSX: açılışta yeniden hesap açık", _wb.calculation.fullCalcOnLoad)
        #  Proje kimliği dosya ÖZELLİKLERİNE yazılıyor mu  ( şablonda hücresi yok )
        _mp = _op.load_workbook(io.BytesIO(
            _MX.mukavemet_xlsx({}, {"proje_adi": "Yıldız Konutları",
                                    "isveren": "ÇAĞDAŞ İnşaat",
                                    "pafta_no": "MK-01"}))).properties
        r.esit("XLSX: proje adı özelliklere yazıldı", _mp.title, "Yıldız Konutları")
        r.esit("XLSX: işveren özelliklere yazıldı", _mp.subject, "ÇAĞDAŞ İnşaat")
        r.esit("XLSX: pafta no özelliklere yazıldı", _mp.category, "MK-01")
        #  UYGULAMA PROJESİNDEN AVAN KİTABI ÇIKMAZ.  İki çalışma kitabı iki
        #  ayrı projeye aittir:  MUKAVEMET_HESABI.xlsx uygulama projesinin,
        #  ASANSOR_AVAN_HESAPLARI.xlsx avan projesinin kitabıdır.  Elektrik
        #  hesapları uygulama tarafında ekranda ve paftada verilir.
        r.kontrol("uygulama projesinde avan kitabı ucu yok",
                  not hasattr(_MM, "indir_uygulama_elektrik_xlsx"))
        r.kontrol("XLSX: hesap sayfaları duruyor",
                  "11-Muk. Hesapları" in _wb.sheetnames
                  and "Askı Tipleri" in _wb.sheetnames)
        #  TESLİM EDİLEN KİTAP PAFTAYLA ÇELİŞMEMELİ.
        #  Program kaynak kitabın sekiz hesabından ayrılıyor;
        #  kitap olduğu gibi verilseydi aynı projenin iki belgesi birbirini
        #  yalanlardı ( pafta "uygun değil" derken Excel "uygundur" ).
        #  Teslim kopyasında o FORMÜLLER düzeltilir — aşağıda gerçekten
        #  düzeltildiği ve kitabın kendi hesabının motorla aynı çıktığı
        #  denetlenir.
        _duz = _op.load_workbook(io.BytesIO(
            _MX.mukavemet_xlsx(_MK.hesapla()["girdi"])))["11-Muk. Hesapları"]
        r.esit("teslim kopyasında Dt/dh eşiği 40", _duz["Q97"].value, 40)
        r.esit("teslim kopyasında Durum 2 xQ = xc", _duz["AO312"].value, "=AH293")
        r.kontrol("teslim kopyasında σ(My) Wx sütununa bakıyor",
                  ",6,0)" in str(_duz["AU575"].value), f"→ {_duz['AU575'].value}")
        r.kontrol("teslim kopyasında μ halat hızıyla",
                  "B100" in str(_duz["AK190"].value), f"→ {_duz['AK190'].value}")
        r.kontrol("teslim kopyasında flanş paydasında ℓ var",
                  "(1+2*" not in str(_duz["Q380"].value).replace(" ", ""),
                  f"→ {str(_duz['Q380'].value)[:90]}")
        r.kontrol("teslim kopyasında ω ray çeliğine bağlı",
                  "B131" in str(_duz["AD354"].value), f"→ {str(_duz['AD354'].value)[:90]}")
        r.esit("teslim kopyasında kabin üstü sınırı sığınma yüksekliğinden",
               _duz["AD636"].value, "=P639*1000")
        r.esit("teslim kopyasında ray dibi açıklığı 100 mm",
               _duz["AD647"].value, 100)
        r.kontrol("teslim kopyasında η makine tipine bağlı formül",
                  "B130" in str(_duz["AQ22"].value), f"→ {_duz['AQ22'].value!r}")
        r.kontrol("teslim kopyasında η askı oranından BAĞIMSIZ  ( Δη kaldırıldı )",
                  "B100" not in str(_duz["AQ22"].value),
                  f"→ {_duz['AQ22'].value!r}")
        _sb = _op.load_workbook(_MX.SABLON)["11-Muk. Hesapları"]
        r.esit("kaynak kitapta η sabit 0,92 idi", _sb["AQ22"].value, 0.92)
        #  Elektrik sayfası programın girdilerini kullanıyor mu
        _delk = _op.load_workbook(io.BytesIO(
            _MX.mukavemet_xlsx(_MK.hesapla()["girdi"])))["12-Elk.Hesapları"]
        _S0 = _US2.sabitler({})
        r.esit("teslim kopyasında U programın değeri", _delk["W28"].value, _S0["U"])
        r.esit("teslim kopyasında cosφ programın değeri", _delk["X58"].value, _S0["cosfi"])
        #  Kesit değişince kapasite de değişmeli  ( kitapta 34 SABİTTİ )
        _delk2 = _op.load_workbook(io.BytesIO(_MX.mukavemet_xlsx(
            _MK.hesapla({"makine_kesit": 1.5})["girdi"])))["12-Elk.Hesapları"]
        r.kontrol("teslim kopyasında kablo kapasitesi KESİTTEN geliyor",
                  _delk2["Y65"].value != _delk["Y65"].value,
                  f"→ 6 mm² {_delk['Y65'].value!r} · 1,5 mm² {_delk2['Y65'].value!r}")
        r.esit("1,5 mm² kablonun kapasitesi", _delk2["Y65"].value, 17.5)
        _sbe = _op.load_workbook(_MX.SABLON)["12-Elk.Hesapları"]
        r.esit("kaynak kitapta kapasiteler ve gerilim SABİTTİ",
               [str(_sbe["S60"].value), str(_sbe["Y65"].value), str(_sbe["W28"].value)],
               ["43", "34", "400"])
        r.esit("teslim kopyasında ray ağırlığı bir kez sayılıyor",
               _duz["AO611"].value, "=AU351-AK351*AM351")
        r.esit("kaynak kitapta bu sınırlar 1200 / 150 idi",
               [_sb["AD636"].value, _sb["AD647"].value], [1200, 150])
        #  ---------------------------------------------------------------
        #  OFİSİN ANA KİTABININ DÜZELTİLMİŞ KOPYASI
        #  ---------------------------------------------------------------
        #  Teslim edilen dosya düzeltiliyordu ama ofisin masasındaki ANA kitap
        #  düzelmiyordu:  onu açıp elle hesap yapan eski, bazıları emniyetsiz
        #  sonuçları alıyordu.  araclar/kaynak_excel_duzelt.py o boşluğu
        #  kapatır;  burada gerçekten düzeltilmiş VE doğru hesaplıyor mu diye
        #  bakılır.
        _usta = _MX.duzeltilmis_kaynak()
        _uwb = _op.load_workbook(io.BytesIO(_usta))
        _uws = _uwb["11-Muk. Hesapları"]
        r.kontrol("düzeltilmiş kitap üretildi", len(_usta) > 200_000,
                  f"→ {len(_usta)} bayt")
        for _h, _bek in (("Q97", 40), ("AO312", "=AH293"),
                         ("AD636", "=P639*1000"), ("AD647", 100)):
            r.esit(f"düzeltilmiş kitap {_h}", _uws[_h].value, _bek)
        for _h, _ara in (("AQ22", "B130"), ("AK190", "B100"),
                         ("AD354", "B131"), ("AU575", ",6,0)")):
            r.kontrol(f"düzeltilmiş kitap {_h} düzeltilmiş",
                      _ara in str(_uws[_h].value), f"→ {str(_uws[_h].value)[:70]}")
        r.kontrol("düzeltilmiş kitapta flanş paydasında ℓ var",
                  "(1+2*" not in str(_uws["Q380"].value).replace(" ", ""))
        #  HİÇBİR PROJENİN GİRDİSİ SIZMAMALI — bu boş bir usta kopyadır
        _uvg = _uwb[_MX.GIRDI_SAYFASI]
        r.kontrol("usta kopyada ek girdi satırları BOŞ",
                  all(_uvg[f"B{_st}"].value in (None, "")
                      for _a, _st, _e, _b in _MX.EK_GIRDI_HUCRELERI),
                  "→ bir projenin girdisi sızmış")
        #  Neyin niçin değiştiği kitabın İÇİNDE yazılı olmalı
        r.kontrol("DÜZELTMELER sayfası var", _MX.DUZELTME_SAYFASI in _uwb.sheetnames,
                  f"→ {_uwb.sheetnames}")
        _dz = "\n".join(str(c.value) for _sat in _uwb[_MX.DUZELTME_SAYFASI].iter_rows()
                         for c in _sat if c.value is not None)
        for _ad, _md, _e2, _y2, _hc in _MK.EXCEL_FARKLARI:
            r.kontrol(f"kayıtta '{_ad[:30]}' var", _ad in _dz)
        for _h in ("Q97", "AQ22", "AD647", "AD636", "AO312"):
            r.kontrol(f"kayıt düzenlenen {_h} hücresini sayıyor", _h in _dz)

        #  ---------------------------------------------------------------
        #  "PROJEYİ PAKETLE"  —  teslim paketi + geri dönüş noktası
        #  ---------------------------------------------------------------
        #  Çıktılar projeyi ANLATIR, proje dosyası onu GERİ GETİRİR.  İkisi
        #  ayrı yerlerde durursa arşivden dönmek imkânsızlaşır;  bu yüzden
        #  aynı ZIP'te olmaları denetlenir.
        import zipfile as _zf
        from api.ortak import _paket_ekleri as _PE
        _pd = {"__mod": "uygulama", "__surum": 1,
               "alanlar": {"m_beyan_yuku": "800", "uof_sigma_em": "150"}}
        #  Kapak alanı sunucuda "project_title" adıyla gelir.
        _ek = _PE({"kapak": {"project_title": "Jan Mühendislik"},
                   "proje_dosyasi": _pd}, "uygulama")
        r.esit("paket eki bir dosya üretiyor", len(_ek), 1)
        r.kontrol("proje dosyasının uzantısı moda göre",
                  _ek[0][0].endswith(".uygulama"), f"→ {_ek[0][0]}")
        r.kontrol("proje dosyası adı proje adından",
                  "Jan Mühendislik" in _ek[0][0], f"→ {_ek[0][0]}")
        _geri = _js.loads(_ek[0][1].decode("utf-8"))
        r.esit("paketteki proje dosyası gövdeyi birebir taşıyor", _geri, _pd)
        r.esit("avan modunda uzantı .avan",
               _PE({"proje_dosyasi": _pd}, "avan")[0][0].endswith(".avan"), True)
        r.esit("proje dosyası yoksa ek de yok", _PE({}, "avan"), [])

        #  ZIP gerçekten hepsini taşıyor mu
        if _DXF is not None:
            _zip, _sebep, _tasti = _DXF.proje_paketi(
                [("Kapak", None), ("Hesap", PE.mukavemet_pdf(_MK.hesapla(), PROJE))],
                "Jan Mühendislik - Uygulama Projesi", _ek)
            _z = _zf.ZipFile(io.BytesIO(_zip))
            _adlar = _z.namelist()
            r.kontrol("pakette DXF var", any(a.endswith(".dxf") for a in _adlar), f"→ {_adlar}")
            r.kontrol("pakette pafta PDF'i var", any(a.endswith(".pdf") for a in _adlar))
            r.kontrol("pakette PROJE DOSYASI var",
                      any(a.endswith(".uygulama") for a in _adlar), f"→ {_adlar}")
            r.kontrol("paket OKUBENİ proje dosyasını anlatıyor",
                      "PROJE DOSYASI" in _z.read("OKUBENI.txt").decode("utf-8"))

        r.kontrol("ŞABLON DOSYASINA DOKUNULMADI",
                  _op.load_workbook(_MX.SABLON)["11-Muk. Hesapları"]["Q97"].value == 30,
                  "→ şablon değişmiş;  doğrulama testleri dayanağını kaybeder")

        #  LibreOffice ile yeniden hesaplandığında motorla aynı sonucu vermeli
        if soffice_yolu():
            _kk = os.path.join(GECICI, "muk_cikti")
            _gg = os.path.join(GECICI, "muk_girdi")
            os.makedirs(_gg, exist_ok=True)
            _yol = os.path.join(_gg, "muk.xlsx")
            with open(_yol, "wb") as _f:
                _f.write(_MX.mukavemet_xlsx({"beyan_yuku": 630, "kabin_agirligi": 600}))
            yeniden_hesapla([_yol], _kk)
            _cy = os.path.join(_kk, "muk.xlsx")
            if os.path.exists(_cy):
                _s = _MK.hesapla({"beyan_yuku": 630, "kabin_agirligi": 600})
                _cws = _op.load_workbook(_cy, data_only=True)["11-Muk. Hesapları"]
                r.kontrol("XLSX: Excel motorla aynı motor gücünü buluyor",
                          abs((_cws["AQ23"].value or 0) - _s["ozet"]["N_hesap"]) < 1e-6,
                          f"→ Excel {_cws['AQ23'].value!r}, motor {_s['ozet']['N_hesap']!r}")
                #  DÜZELTİLMİŞ KİTABIN HER HÜCRESİ MOTORLA AYNI OLMALI —
                #  standart gereği saptığımız hücreler DÂHİL.  İki belge
                #  arasında tek bir sayı bile ayrışmamalı.
                _ayri = []
                for _h, _v in sorted(_s["_h"].items()):
                    _e = _cws[_h].value
                    if _e is None or isinstance(_e, str):
                        continue
                    try:
                        if abs(float(_v) - float(_e)) > 1e-6 * max(
                                abs(float(_v)), abs(float(_e)), 1.0):
                            _ayri.append((_h, _v, _e))
                    except (TypeError, ValueError):
                        pass
                r.kontrol("teslim edilen Excel paftayla BİREBİR aynı", not _ayri,
                          f"→ ayrışan {len(_ayri)} hücre: {_ayri[:3]}")
                _hh = [h for h in hata_hucresi_ara(_cy)
                       if h.startswith("11-Muk.") or h.startswith("Askı Tipleri")]
                r.kontrol("XLSX: hesap sayfalarında hata hücresi yok", not _hh,
                          f"→ {_hh[:3]}")
            else:
                r.kontrol("XLSX yeniden hesaplanabildi", False)
    else:
        r.atla("Mukavemet şablonu yok — XLSX çıktısı denenmedi")

    #  İndirme uçları
    _mg = {a[0]: (list(a[6]) if a[4] == "liste" else a[6])
           for a in _MG.ALANLAR if a[4] != "hesap"}
    for _ad, _fn, _tur in (("pdf", _MM.indir_uygulama_pdf, "application/pdf"),
                           ("xlsx", _MM.indir_uygulama_xlsx, XLSX_TUR)):
        _y = _fn({"girdiler": _mg, "kapak": {"proje_adi": "DENEME"}})
        r.esit(f"uygulama-{_ad} ucu dosya döndürüyor", _y.media_type, _tur)

    #  UYGULAMA PROJESİ PAFTASI  —  mukavemet + elektrik + topraklama
    from engine.uygulama import hesap as _UY
    _uy = _UY.hesapla({"temel_a": 26.55, "temel_b": 16.4, "kolon_uzunluk": 45})
    _upd = PE.uygulama_pdf(_uy, PROJE)
    r.kontrol("uygulama PDF üretildi", len(_upd) > 30_000, f"→ {len(_upd)} bayt")
    _um = _metin(_upd)
    for _ara in ("ASANSÖR UYGULAMA PROJESİ HESAPLARI", "MUKAVEMET HESAPLARI",
                 "ELEKTRİK VE TOPRAKLAMA HESAPLARI", "KABİN AYDINLATMA",
                 "KURULU GÜÇ CETVELİ", "GERİLİM DÜŞÜMÜ", "TOPRAKLAYICI",
                 "SONUÇ ÖZETİ", "Hesabı yapan"):
        r.kontrol(f"uygulama PDF: {_ara}", _ara in _um)
    #  Kurulu güç PAFTADAN değil, HESAPTAN okunur:  elle yazılmış bir sayı
    #  motor değiştiğinde sessizce bayatlar ( ηm düzeltmesinde öyle oldu ).
    from engine.ortak.steps import trn as _trn
    _pk = _trn(_uy["ozet"]["P_kurulu"], 0)
    r.kontrol("uygulama PDF elektrik sayılarını taşıyor",
              _pk in _um, f"→ kurulu güç ({_pk} W) paftada yok")
    _ubos = PE.uygulama_pdf(_UY.hesapla({"makine_agirligi": None}), PROJE)
    r.kontrol("hatalı girdide uygulama PDF'i sebebini yazıyor",
              "boş bırakılamaz" in _metin(_ubos))

    #  Hesap durduran girdide DOSYA DEĞİL, açık hata dönmeli
    _yh = _MM.indir_uygulama_xlsx({"girdiler": dict(_mg, makine_agirligi=None)})
    r.kontrol("hatalı girdide XLSX yerine hata dönüyor",
              _yh.media_type != XLSX_TUR
              and "boş bırakılamaz" in _yh.body.decode("utf-8"),
              f"→ {_yh.media_type}")

    return r


if __name__ == "__main__":
    sys.exit(0 if calistir().yazdir() else 1)

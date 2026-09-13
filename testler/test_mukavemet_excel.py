# -*- coding: utf-8 -*-
"""
TEST 10  —  MUKAVEMET ↔ EXCEL UYUMU        ( uygulama projesi )

TEST 9 tek senaryoyu ( Excel'in kendi örneğini ) denetler.  Bu test GİRDİ
UZAYINI denetler:  onlarca farklı girdi bileşimi kaynak çalışma kitabının
"Veri Girişi" sayfasına yazılır, LibreOffice bütün formülleri yeniden
hesaplar ve motorun ürettiği her değer hücre hücre karşılaştırılır.

Formülü değil, formülün BÜTÜN GİRDİLERDEKİ davranışını kilitler:  ray
profili, halat çapı, kanal şekli/işlemesi, güvenlik tertibatı tipi, askı
oranı, karşı ağırlık yeri, durak sayısı, hız, yük …

LibreOffice kurulu değilse test ATLANIR.
"""
import os
import shutil
import sys
import warnings

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                          # noqa: E402

from engine.uygulama import mukavemet as MK                       # noqa: E402
from exports import mukavemet_xlsx as MX                          # noqa: E402
from engine.uygulama import mukavemet_girdi as MG                 # noqa: E402
from engine.uygulama import mukavemet_tablolari as MT             # noqa: E402
from testler.ortak import (Rapor, hata_hucresi_ara,      # noqa: E402
                           soffice_yolu, yeniden_hesapla)

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SABLON = os.path.join(KOK, "templates", "MUKAVEMET_HESABI.xlsx")
GECICI = os.path.join(KOK, ".gecici_mukavemet")
SAYFA = "11-Muk. Hesapları"

#  KAYNAK EXCEL'DEN BİLEREK AYRILAN HÜCRELER — karşılaştırma dışı.
#  Liste MOTORDAN okunur ( engine.mukavemet.EXCEL_FARKLARI );  testin kendi
#  kopyasını tutması, sapmanın gerekçesiyle testin ayrışmasına yol açardı.
#  Uyulması gereken standart TS EN 81-20 / TS EN 81-50'dir:  bu hücrelerde
#  Excel standarttan sapıyor, motor standardı uyguluyor.
AYRILAN = set(MK.FARKLI_HUCRELER)


def senaryolar():
    """Girdi uzayını tarayan senaryolar  ( ad , girdi sözlüğü )."""
    v = MG.varsayilanlar()
    s = []

    def E(ad, **kw):
        g = dict(v)
        g.update(kw)
        s.append((ad, g))

    def durak(n, h=3000, son=3750):
        dy = [h] * (n - 1) + [son]
        return {"durak_yukseklikleri": dy, "son_kat_yuksekligi": son,
                "seyir_mesafesi": (n - 1) * h / 1000.0}

    E("varsayılan")

    #  SARILMA AÇISI zorunlu girdidir ve varsayılanı YOKTUR:  öteki bütün
    #  senaryolarda tahrik sınırları hesaplanmaz ( HESAP EKSİK ).  Kitap ise
    #  açıyı kendi şase ölçülerinden türetir ( 137,89° ).  Bu iki senaryo
    #  sınır hücrelerinin ( S184 · AA184 · O242 … O285 ) motorda gerçekten
    #  üretildiğini ve kitaptan BİLİNÇLİ olarak ayrıştığını gösterir.
    E("sarılma açısı 180°", sarilma_acisi=180)
    E("sarılma açısı 150°", sarilma_acisi=150)

    #  DENGE ZİNCİRİ:  P'ye giren MCR yolunu açar  ( m.5.2.1.8.5 · m.5.7.2.3.2 ).
    #  Zincirsiz senaryolarda MCR = 0'dır ve ray / kuyu tabanı sapmaları
    #  görünmez;  bu senaryo onları sınar.
    E("denge zinciri var", denge_zinciri="Var")
    #  MAKİNE RAYLARA BİNİYOR  ( MRL — m.5.7.2.3.7 ):  Maux artık 150 N değil.
    #  Seçim yalnız MRL'de uygulanır ( makine dairesi varsa makine kendi
    #  kaidesindedir ve yük iki kez sayılmamalıdır ).
    E("makine raylara biniyor", mk_yok=True, makine_raya_biniyor=MT.MAKINE_YUK_YOLU_RAY)
    #  Zincir + karşı ağırlıkta güvenlik tertibatı:  ağırlık rayının ve kuyu
    #  tabanının ( FAR ) zinciri gördüğü tek bileşim.
    E("zincir + ağırlık güvenlik tertibatı", denge_zinciri="Var",
      agirlik_guvenlik_tertibati=MT.DARBE_TIPLERI_ADLARI[0])
    #  Regülatör halatı katalog verisi:  tabloyu ezer  ( AI134 · AI136 ).
    E("regülatör katalog halatı", reg_halat_kopma_kN=28,
      reg_halat_birim_kutle=0.30)

    #  Ray profilleri  —  kabin ve karşı ağırlık ayrı ayrı
    for p in MT.RAY_PROFILLERI:
        E(f"kabin rayı {p}", kabin_ray_profili=p)
        E(f"ağırlık rayı {p}", agirlik_ray_profili=p)

    #  Halat çapları  ( askı ) ve regülatör halatı
    for c in MT.HALAT_CAPLARI:
        E(f"halat Ø{c}", halat_capi=c)
    for c in (6, 6.5, 8):
        E(f"regülatör halatı Ø{c}", reg_halat_capi=c)

    #  Kanal şekli / işlemesi  →  Nequiv(t) ve sürtünme faktörü
    for k in MT.KANAL_SEKILLERI:
        #  Düz V kanal sertleştirilmemiş olamaz ( TS EN 81-50 m.5.11.2.3.1.2 ) —
        #  girdide reddedilir;  o kanal sertleştirilmiş işlemeyle sınanır.
        if MT.kanal_turu(k) == "V":
            E(f"kanal {k}", kanal_sekli=k, kanal_isleme="Sertleştirilmiş")
        else:
            E(f"kanal {k}", kanal_sekli=k)
    for k in MT.KANAL_ISLEME_SEKILLERI:
        E(f"kanal işleme {k}", kanal_isleme=k)

    #  Güvenlik tertibatı  →  k1
    for t in MT.DARBE_TIPLERI_ADLARI:
        E(f"güv. tertibatı {t}", guvenlik_tertibati=t)

    #  Karşı ağırlık malzemesi.  DERİNLİK DE VERİLİR:  motor artık onu
    #  malzemeden türetmiyor ( ölçü imal edilen çerçevenin özelliğidir,
    #  bkz. TS EN 81-50 Ek C.2.2 ) ama ÖZGÜN KİTAP türetmeye devam ediyor.
    #  Karşılaştırmanın anlamlı kalması için ikisi aynı sayıyı kullanmalı;
    #  yoksa sınanan şey kitap değil, kaldırdığımız türetme olurdu.
    _KITAP_DERINLIK = {r[0]: r[1] for r in MT.AGIRLIK_MALZEMESI}
    for m in MT.AGIRLIK_MALZEMELERI:
        E(f"ağırlık malzemesi {m}", agirlik_malzemesi=m,
          agirlik_derinligi=_KITAP_DERINLIK[m])
    #  RAY ARASI ARTIK SERBEST ÖLÇÜ.  Üç tablo noktası ( 700 · 1050 · 1400 )
    #  birebir korunmalı;  ARADAKİ ve TABLO DIŞI değerlerde de LibreOffice
    #  motorla aynı genişliği bulmalı, yoksa kitap #YOK der ve pafta ile
    #  ayrışır.  İmalatçı genişliği girilen durum da denenir — o zaman tablo
    #  hiç kullanılmaz.
    #  KARŞI AĞIRLIĞIN ÖLÇÜLERİ BU DÖNGÜDE DEĞİŞTİRİLMEZ.  Buradaki
    #  senaryolar ÖZGÜN kitaba yazılır ve özgün kitapta bu iki girdinin
    #  hücresi YOKTUR — ölçüleri hâlâ kendi tablolarından türetir
    #  ( VLOOKUP B118 → derinlik, HLOOKUP B119 → genişlik ).  Motor ise
    #  girilen ölçüyü kullanır;  ikisini farklı değerlerle karşılaştırmak
    #  özgün kitabı değil, kaldırdığımız türetmeyi sınamak olurdu.
    #  Serbest ölçüler TESLİM EDİLEN kitapta denetlenir ( aşağıda ).

    #  Ray çeliği  →  σperm
    for rm in MT.RAY_CELIKLERI:
        E(f"ray çeliği Rm {rm}", ray_celigi_rm=rm)

    #  Askı oranı  →  halat boyu, T1/T2, güvenlik katsayısı
    E("askı 1:1", aski_orani=1)

    #  Karşı ağırlık yeri  →  halat arası
    E("ağırlık solda", agirlik_yeri="Sol")
    E("ağırlık arkada", agirlik_yeri="Arka")

    #  Hız  →  μ frenleme, serbest boşluk
    for h in (0.63, 1.6, 2.5, 4):
        E(f"hız {h} m/s", beyan_hizi=h)

    #  Yük / kabin  →  kabin alanı, Fs, T1/T2
    E("630 kg küçük kabin", beyan_yuku=630, kabin_genisligi=1100,
      kabin_derinligi=1400, kabin_agirligi=600)
    E("1600 kg büyük kabin", beyan_yuku=1600, kabin_genisligi=1900,
      kabin_derinligi=2000, kabin_agirligi=1200, kabin_ray_profili="125 x 82 x 16")
    E("2000 kg", beyan_yuku=2000, kabin_genisligi=2000, kabin_derinligi=2200,
      kabin_agirligi=1400, kabin_ray_profili="127 x 89 x 16")
    E("320 kg", beyan_yuku=320, kabin_genisligi=900, kabin_derinligi=1000,
      kabin_agirligi=450)

    #  Durak sayısı  →  ray boyu, Mg, MSR, sığınma
    for n in (2, 5, 12, 20):
        E(f"{n} durak", **durak(n))
    E("yüksek son kat", **durak(8, son=5000))

    #  NPU kesitleri  →  makine konstrüksiyonu
    for npu in (80, 160, 200):
        E(f"NPU {npu}", dikine_kiris=npu, yan_yatak=npu)

    #  Acil frenleme yavaşlaması  →  T1/T2
    #  a = gn ( 9,81 ) sınır hâlinde T1 sıfırlanır ve Excel #SAYI/0! üretir;
    #  o uç TEST 9'da motor tarafında denenir, burada karşılaştırılamaz.
    #  0,5 m/s²  TS EN 81-50 m.5.11.2.2.2'nin ALT SINIRIDIR;  altı artık
    #  reddedilir ( bkz. EXCEL_FARKLARI ), o yüzden sınırın kendisinden başlanır.
    for a in (0.5, 1.5, 5.0):
        E(f"a = {a} m/s²", acil_frenleme_a=a)

    #  Bükülgen kablo  →  MTrav
    #  Yalnız 2. kablo tipi taranır:  Excel 1. kabloyu VLOOKUP ile aramaz,
    #  TABLOLAR!G62'ye ( 24 x 0,75 ) sabitler — bilinen kaynak hatası, motor
    #  ikisini de arar.  1. kablonun etkisi TEST 9'da denetlenir.
    for k in MT.KABLO_TIPLERI:
        E(f"kablo {k}", kablo_tipi_2=k)

    #  Pervaz  →  kabin alanı dalı
    E("uzun pervaz 150", uzun_pervaz=150)

    #  Geometri  →  xc · xp · Dxa · sığınma
    E("ray-kapı arası 500", ray_kapi_arasi=500)
    #  RAY EKSENİ KABİN MERKEZİNİ GEÇİYOR  →  xc ≤ 0, kuvvetler NEGATİFE düşer.
    #  Bu bölge kapsam dışıydı ve iki hatayı gizliyordu:  kitabın σ/δ
    #  hücreleri işaretli hesaplıyordu ( negatif sehim sessizce geçiyordu ) ve
    #  kapı konumu xi ham mesafe olarak yazılıyordu.
    E("ray-kapı arası 830  ( xc = 0 )", ray_kapi_arasi=830)
    E("ray-kapı arası 1200 ( xc < 0 )", ray_kapi_arasi=1200)
    E("kabin kaçıklığı 60", kabin_kaciklik=60)
    E("ağır kabin kapısı", kapi_agirligi=180, kapi_mekanizma_payi=90)
    E("dar konsol aralığı", kabin_konsol_arasi=2000, agirlik_konsol_arasi=2000)
    E("geniş paten arası", kabin_paten_arasi=4200, agirlik_paten_arasi=4200)
    return s


def _yaz(g, yol):
    """Girdileri kaynak çalışma kitabının 'Veri Girişi' sayfasına yazar."""
    wb = openpyxl.load_workbook(SABLON)
    ws = wb["Veri Girişi"]
    for anahtar, hucre, _e, _b, tur, _s, _v in MG.ALANLAR:
        if tur == "hesap":
            continue                     # bunlar Excel'de formüldür, ezilmez
        #  Kaynak Excel'de KARŞILIĞI OLMAYAN alanlar ( ör. paten balata boyu )
        #  boş hücre adresi taşır — yazılacak/okunacak yerleri yoktur.
        if not hucre:
            continue
        if tur == "liste":
            dy = g[anahtar]
            for i, h in enumerate(MG.DURAK_HUCRELERI):
                ws[h] = dy[i] if i < len(dy) else None
            continue
        ws[hucre] = g[anahtar]
    #  T75/B KAYNAK KİTAPTA YOKTUR.  Motorun kataloğu ISO 7465'ten bir satır
    #  genişletildi;  ham kitap o rayı tanımadığı için bütün VLOOKUP'lar #YOK
    #  verir ve senaryo HİÇ kıyaslanamaz olurdu.  Teslim edilen kopyaya
    #  uygulanan genişletmenin AYNISI buraya da uygulanır — satır EKLER,
    #  kitabın hiçbir mevcut değerini değiştirmez  ( bkz.
    #  exports/mukavemet_xlsx._ray_tablosu_genislet ).  Böylece bu test
    #  genişletmenin kendisini de doğrulamış olur.
    MX._ray_tablosu_genislet(wb)
    wb.save(yol)


def calistir():
    print("\n\033[1mTEST 10 — MUKAVEMET ↔ EXCEL UYUMU\033[0m"
          "   (girdiler şablona yazılır, LibreOffice yeniden hesaplar)")
    r = Rapor("Mukavemet ↔ Excel")
    if not os.path.isfile(SABLON):
        r.atla(f"Kaynak Excel yok — {os.path.basename(SABLON)}")
        return r
    if not soffice_yolu():
        r.atla("LibreOffice bulunamadı — bu test yalnız LibreOffice kurulu "
               "makinede çalışır.")
        return r
    warnings.filterwarnings("ignore")

    shutil.rmtree(GECICI, ignore_errors=True)
    giris, cikis = os.path.join(GECICI, "girdi"), os.path.join(GECICI, "cikti")
    os.makedirs(giris, exist_ok=True)

    sen = senaryolar()
    ayrisan = set()
    sonuclar, dosyalar = [], []
    for i, (ad, g) in enumerate(sen):
        s = MK.hesapla(g)
        if not s["aktif"]:
            r.kontrol(f"[{ad}] senaryo geçerli", False, f"→ {s['hata']}")
            continue
        yol = os.path.join(giris, f"m{i:03d}.xlsx")
        _yaz(s["girdi"], yol)
        dosyalar.append(yol)
        sonuclar.append((i, ad, s))

    print(f"   {len(dosyalar)} çalışma kitabı üretildi — LibreOffice yeniden "
          "hesaplıyor…")
    yeniden_hesapla(dosyalar, cikis)

    for i, ad, s in sonuclar:
        yol = os.path.join(cikis, f"m{i:03d}.xlsx")
        if not os.path.exists(yol):
            r.kontrol(f"[{ad}] dosya üretilemedi", False)
            continue
        ws = openpyxl.load_workbook(yol, data_only=True)[SAYFA]
        for hucre, deger in sorted(s["_h"].items()):
            if hucre in AYRILAN:
                continue
            bek = ws[hucre].value
            r.esit(f"[{ad}] {SAYFA}!{hucre}", deger, bek)
        #  Sapmalar gerçekten ayrışıyor mu — sessizce Excel'e dönmesinler.
        #  Bazı senaryolarda fark doğal olarak sıfırlanır ( ör. kabin merkezi
        #  ray ekseninde ise Durum 2'nin xQ'su zaten 0'dır ), o yüzden
        #  "en az bir senaryoda ayrışıyor" aranır, her senaryoda değil.
        for hucre in AYRILAN:
            if (hucre in s["_h"] and ws[hucre].value is not None
                    and not _esit(s["_h"][hucre], ws[hucre].value)):
                ayrisan.add(hucre)
        for e in hata_hucresi_ara(yol):
            #  Kaynak dosya, mukavemetle ilgisi olmayan sayfalarda da hata
            #  hücresi taşıyabilir;  yalnız hesap sayfalarını denetliyoruz.
            if e.startswith(SAYFA) or e.startswith("Askı Tipleri"):
                r.kontrol(f"[{ad}] Excel hata hücresi", False, e)

    #  Her sapma senaryoların EN AZ BİRİNDE ayrışmalı;  yoksa ya kod Excel'e
    #  geri dönmüş ya da EXCEL_FARKLARI bayatlamıştır.
    for hucre in sorted(AYRILAN):
        r.kontrol(f"sapma {hucre} en az bir senaryoda Excel'den ayrışıyor",
                  hucre in ayrisan,
                  "→ hiçbir senaryoda ayrışmadı; motor Excel'e dönmüş olabilir "
                  "ya da EXCEL_FARKLARI bayatlamış")
    #  ---------------------------------------------------------------
    #  OFİSİN DÜZELTİLMİŞ ANA KİTABI KENDİ BAŞINA DOĞRU HESAPLIYOR MU
    #  ---------------------------------------------------------------
    #  Yukarıdaki senaryolar ÖZGÜN kitabı denetler ( sapmalarımızın dayanağı
    #  odur ).  Burada ise araclar/kaynak_excel_duzelt.py'nin ürettiği kitap
    #  LibreOffice'e yeniden hesaplatılır:  ofis onu tek başına açıp
    #  kullandığında pafta ile aynı sonucu vermelidir.
    _uk = os.path.join(GECICI, "usta")
    _ug = os.path.join(_uk, "in")
    os.makedirs(_ug, exist_ok=True)
    with open(os.path.join(_ug, "usta.xlsx"), "wb") as _f:
        _f.write(MX.duzeltilmis_kaynak())
    yeniden_hesapla([os.path.join(_ug, "usta.xlsx")], os.path.join(_uk, "out"))
    _uyol = os.path.join(_uk, "out", "usta.xlsx")
    if not os.path.isfile(_uyol):
        r.kontrol("düzeltilmiş kitap LibreOffice ile açıldı", False,
                  "→ dönüştürme başarısız")
    else:
        _uws = openpyxl.load_workbook(_uyol, data_only=True)[SAYFA]
        _s = MK.hesapla()
        _tut = 0
        for _h, _m in sorted(_s["_h"].items()):
            _v = _uws[_h].value
            if not isinstance(_v, (int, float)) or not isinstance(_m, (int, float)):
                continue
            if isinstance(_v, bool) or isinstance(_m, bool):
                continue
            _tut += 1
            r.kontrol(f"[düzeltilmiş kitap] {SAYFA}!{_h}", _esit(_m, _v),
                      f"→ motor {_m!r}, kitap {_v!r}")
        r.kontrol("düzeltilmiş kitapta karşılaştırılan hücre var", _tut > 100,
                  f"→ yalnız {_tut} hücre")
        #  Kitabın kendi hesabı, düzeltilmiş eşikleri de göstermeli
        #  AQ22:  Δη KALDIRILDIĞI için 2:1 askıda da ofis tablosunun kendisi
        #  ( dişlisiz 0,85 ) geçerlidir — eskiden 0,85 − 0,10 = 0,75 idi.
        for _h, _bek in (("Q97", 40), ("AD636", 1000), ("AD647", 100),
                         ("AQ22", 0.85)):
            r.kontrol(f"[düzeltilmiş kitap] {_h} = {_bek}",
                      _esit(_uws[_h].value, _bek), f"→ {_uws[_h].value!r}")
        for _e in hata_hucresi_ara(_uyol):
            if _e.startswith(SAYFA) or _e.startswith("Askı Tipleri"):
                r.kontrol("[düzeltilmiş kitap] Excel hata hücresi", False, _e)

    #  ---------------------------------------------------------------
    #  TESLİM EDİLEN KİTAP  —  KARŞI AĞIRLIĞIN SERBEST ÖLÇÜLERİ
    #  ---------------------------------------------------------------
    #  Özgün kitap bu iki ölçüyü tablodan türetiyordu ( ray arası → genişlik
    #  üç değere kilitliydi ).  TS EN 81-50 Ek C.2.2 ikisini de VERİ olarak
    #  ister;  türetme kaldırıldı.  Teslim edilen kitap girilen ölçüleri
    #  kullanmalı, yoksa pafta ile kitap ayrışır ve mühendis kitabı açtığında
    #  başka bir gerilme görür.  Tablo DIŞI değerler seçildi.
    _sk = os.path.join(GECICI, "serbest")
    _sg = os.path.join(_sk, "in")
    os.makedirs(_sg, exist_ok=True)
    _olcu = {"agirlik_genisligi": 850, "agirlik_derinligi": 130,
             "agirlik_ray_arasi": 1200}
    try:
        _bayt = MX.mukavemet_xlsx(dict(MG.varsayilanlar(), **_olcu), None)
    except FileNotFoundError:
        _bayt = None
    if _bayt is None:
        r.atla("mukavemet şablonu yok — serbest ölçü denetimi atlandı")
    else:
        with open(os.path.join(_sg, "serbest.xlsx"), "wb") as _f:
            _f.write(_bayt)
        yeniden_hesapla([os.path.join(_sg, "serbest.xlsx")],
                        os.path.join(_sk, "out"))
        _syol = os.path.join(_sk, "out", "serbest.xlsx")
        if not os.path.isfile(_syol):
            r.kontrol("[serbest ölçü] kitap LibreOffice ile açıldı", False,
                      "→ dönüştürme başarısız")
        else:
            _sws = openpyxl.load_workbook(_syol, data_only=True)[SAYFA]
            #  Girilen ölçüler kitabın kendi hücrelerine geçmiş mi
            r.esit("[serbest ölçü] Gy kitapta", _sws["AH551"].value, 850)
            r.esit("[serbest ölçü] Gx kitapta", _sws["AH550"].value, 130)
            #  Ve ASIL KONTROL:  kitabın bütün hesabı motorla aynı mı
            _ss = MK.hesapla(_olcu)
            _st = 0
            for _h, _m in sorted(_ss["_h"].items()):
                _v = _sws[_h].value
                if (not isinstance(_v, (int, float))
                        or not isinstance(_m, (int, float))
                        or isinstance(_v, bool) or isinstance(_m, bool)):
                    continue
                _st += 1
                r.kontrol(f"[serbest ölçü] {SAYFA}!{_h}", _esit(_m, _v),
                          f"→ motor {_m!r}, kitap {_v!r}")
            r.kontrol("[serbest ölçü] karşılaştırılan hücre var", _st > 100,
                      f"→ yalnız {_st} hücre")
            for _e in hata_hucresi_ara(_syol):
                if _e.startswith(SAYFA):
                    r.kontrol("[serbest ölçü] Excel hata hücresi", False, _e)

    #  ---------------------------------------------------------------
    #  TESLİM EDİLEN KİTAP  —  SÜRTÜNME ÇARPANI f KANAL ŞEKLİNE BAĞLI
    #  ---------------------------------------------------------------
    #  Kitap f'i kanal şeklinden bağımsız hep V kanal bağıntısıyla kuruyor,
    #  γ ve β'yı hücreye çiviliyordu ( 38° · 90° ).  Motor bunu çoktan
    #  düzeltmişti ( EXCEL_FARKLARI ) ama teslim kopyası düzeltilmemişti:
    #  pafta ile kitap tahrik sınırında 10 kanal × işleme birleşiminin
    #  8'inde ayrışıyordu ( ör. yarım daire bloke sınırı 2,19 · 6,89 ).
    #  Yukarıdaki senaryolar HAM kitabı denetler;  burası TESLİM kopyasını.
    #  Ofis açıları değiştirilmiş bir tur da koşulur:  kitap açıları kendi
    #  sabitinden değil, projenin ofis sabitinden okumalı.
    #  DOSYA ADLARI BENZERSİZ:  bir ön denetim adı işlemenin ilk beş harfiyle
    #  kurmuş ( "Sertl" ), iki işleme aynı dosyaya yazılıp birbirini ezmişti.
    _fk = os.path.join(GECICI, "kanal_f")
    _fg = os.path.join(_fk, "in")
    os.makedirs(_fg, exist_ok=True)
    _F_HUC = ("AU206", "AV211", "AJ198", "AL202", "AE216",
              "O242", "O257", "O271", "O285")
    _OFIS_DEGISIK = {"kanal_gama_v": 45, "kanal_gama_yd": 30, "kanal_beta": 100}
    _fsen = []
    for _ofis in (None, _OFIS_DEGISIK):
        for _k in MT.KANAL_SEKILLERI:
            for _i in MT.KANAL_ISLEME_SEKILLERI:
                #  Sertleştirilmemiş düz V standart dışıdır ve motor onu reddeder;
                #  f karşılaştırması yapılamaz — aşağıda ayrı sınanır.
                if MT.kanal_turu(_k) == "V" and _i == "Sertleştirilmemiş":
                    continue
                #  Açı kanalın sarım sayısına uymalı:  tek sarımda ≤ 180°,
                #  çift sarımda > 180° ( aykırısı girdide reddedilir ).
                _aci = 330 if (MT.kanal_gecis_sayisi(_k) or 1) > 1 else 180
                _g = {"kanal_sekli": _k, "kanal_isleme": _i, "sarilma_acisi": _aci}
                if _ofis:
                    _g["_ofis"] = dict(_ofis)
                _fsen.append((f"f{len(_fsen):02d}", _g))
    #  GEÇERSİZ GİRDİDE KİTAP DA HÜKÜM VERMEZ.  Motor bu üç girdiyi reddeder;
    #  kitap onları hiç denetlemiyor, tek sarım · 300°'de RAPOR'a üç kez
    #  "UYGUNDUR." yazıyordu.  Kitap motordan bağımsız çağrıldığında ( elle
    #  değiştirilmiş hücre ) da hüküm hücreleri "UYGUN DEĞİLDİR" demeli.
    _gecersiz = (
        ("g_tek300", {"sarilma_acisi": 300, "kanal_isleme": "Sertleştirilmiş"},
         "sarım sayısına"),
        ("g_cift180", {"kanal_sekli": "Yarım Daire Kanal (Çift Sarım)",
                       "sarilma_acisi": 180}, "sarım sayısına"),
        ("g_Vsoft", {"kanal_sekli": "V Kanal", "kanal_isleme": "Sertleştirilmemiş",
                     "sarilma_acisi": 180}, "alt kesilmesiz V"))
    #  D/d < 40 BELGESİ.  Kitabın Dt/dh hükmü ( Z97 ) motorla aynı kuralı
    #  izlemeli:  belgesiz 240 / 6,5 ret, belgeli kabul, oran ≥ 40 iken belge
    #  hükmü değiştirmez.  Sf ( T125 ) belgeden etkilenmemeli.
    _KUCUK = {"tahrik_kasnak_capi": 240, "halat_capi": 6.5,
              "saptirma_kasnak_capi": 240, "kasnak_tek_yon": 2}
    _belgeli = (("d_yok", dict(_KUCUK, kasnak_belgesi="Yok")),
                ("d_var", dict(_KUCUK, kasnak_belgesi="Var")),
                ("d_var40", {"tahrik_kasnak_capi": 320, "halat_capi": 8,
                             "kasnak_belgesi": "Var"}))
    #  TAHRİK HÜCRELERİ TESLİM KİTABINDA MOTORLA AYNI.  Bu hücreler kaynak
    #  kitaptan BİLEREK ayrışır ( AYRILAN ) ve yukarıdaki ana döngü onları
    #  atlar;  teslim kitabının onları motorla aynı kurduğunu hiçbir şey
    #  görmüyordu.  Görünmeyen üç ayrışma buradan çıktı:  denge zinciri
    #  kitapta tahrike girmiyordu, bloke halat kütlesi katalogdan değil
    #  tablodan okunuyordu, 1. gezici kablo 24 × 0,75'e çiviliydi — ve
    #  sürtünme iki kez r'ye bölünüyordu.  q ≠ 0,50'de Gmax da ayrışıyordu.
    _TAHRIK_H = ("AF235", "AJ240", "K242", "AF250", "AJ255", "K257",
                 "AH264", "AF269", "K271", "AH278", "AF283", "K285", "AQ9")
    _tahrik = (
        ("t_11", {"aski_orani": 1, "sarilma_acisi": 150, "kanal_sekli": "V Kanal",
                  "kanal_isleme": "Sertleştirilmiş"}),
        ("t_zincir", {"aski_orani": 2, "sarilma_acisi": 180, "denge_zinciri": "Var"}),
        ("t_katalog", {"aski_orani": 2, "sarilma_acisi": 180, "denge_zinciri": "Var",
                       "halat_birim_kutle": 0.21,
                       "_ofis": {"denge_zinciri_orani": 70, "q_denge": 0.45}}),
        ("t_kablo", {"aski_orani": 2, "sarilma_acisi": 180, "kablo_birim_kutle": 0.44}),
        ("t_kablotip", {"aski_orani": 1, "sarilma_acisi": 170, "kablo_tipi_1": "12 x 1,00",
                        "kat_kapisi_tipi": "Manuel Sağ"}),
        ("t_fr0", {"aski_orani": 2, "sarilma_acisi": 180,
                   "_ofis": {"kuyu_surtunme_kabin": 0, "kuyu_surtunme_agirlik": 0}}),
        ("t_fr3", {"aski_orani": 2, "sarilma_acisi": 170,
                   "_ofis": {"kuyu_surtunme_kabin": 3, "kuyu_surtunme_agirlik": 2.5}}),
    )
    try:
        for _ad, _g in (list(_fsen) + [(a, g) for a, g, _ in _gecersiz] + list(_belgeli)
                        + list(_tahrik)):
            with open(os.path.join(_fg, _ad + ".xlsx"), "wb") as _f:
                _f.write(MX.mukavemet_xlsx(_g))
    except FileNotFoundError:
        _fsen = []
        r.atla("mukavemet şablonu yok — kanal f denetimi atlandı")
    if _fsen:
        yeniden_hesapla([os.path.join(_fg, a + ".xlsx") for a, _ in _fsen]
                        + [os.path.join(_fg, a + ".xlsx") for a, _g, _p in _gecersiz]
                        + [os.path.join(_fg, a + ".xlsx") for a, _g in _belgeli]
                        + [os.path.join(_fg, a + ".xlsx") for a, _g in _tahrik],
                        os.path.join(_fk, "out"))
        for _ad, _g in _tahrik:
            _s = MK.hesapla(_g)
            _tyol = os.path.join(_fk, "out", _ad + ".xlsx")
            if not r.kontrol(f"[tahrik] {_ad} kitabı hesaplandı",
                             _s.get("aktif") and os.path.isfile(_tyol), f"→ {_s.get('hata')}"):
                continue
            _tws = openpyxl.load_workbook(_tyol, data_only=True)[SAYFA]
            for _h in _TAHRIK_H:
                r.kontrol(f"[tahrik] {_ad} · {_h} teslim kitabında motorla aynı",
                          _esit(_s["_h"][_h], _tws[_h].value),
                          f"→ motor {_s['_h'][_h]!r}, kitap {_tws[_h].value!r}")
        for _ad, _g in _belgeli:
            _s = MK.hesapla(_g)
            _b4 = next(b for b in _s["bolumler"] if b["kimlik"] == "aski_halatlari")
            _dd = [a for a in _b4["adimlar"] if str(a.get("aciklama") or "").startswith("Dt / dh =")]
            _motor_uygun = bool(_dd) and not str(_dd[0].get("metin")).startswith("UYGUN DEĞİL")
            _byol = os.path.join(_fk, "out", _ad + ".xlsx")
            if not os.path.isfile(_byol):
                r.kontrol(f"[D/d belgesi] {_ad} kitabı hesaplandı", False)
                continue
            _bws = openpyxl.load_workbook(_byol, data_only=True)[SAYFA]
            _z = str(_bws["Z97"].value or "")
            r.kontrol(f"[D/d belgesi] {_ad} · Z97 motorla aynı hükmü veriyor",
                      _z.startswith("UYGUNDUR") == _motor_uygun,
                      f"→ kitap {_z!r}, motor {_dd[0].get('metin') if _dd else None!r}")
            r.kontrol(f"[D/d belgesi] {_ad} · Sf ( T125 ) motorla aynı",
                      _esit(_s["ozet"]["Sf"], _bws["T125"].value),
                      f"→ motor {_s['ozet']['Sf']!r}, kitap {_bws['T125'].value!r}")
        _z_var = str(openpyxl.load_workbook(os.path.join(_fk, "out", "d_var.xlsx"),
                                            data_only=True)[SAYFA]["Z97"].value or "")
        r.esit("[D/d belgesi] belgeli kitap sapmayı adıyla yazıyor",
               _z_var, MK.BELGEYLE_UYGUN_KITAP)
        r.kontrol("[D/d belgesi] belgesiz 240 / 6,5 kitapta da reddediliyor",
                  str(openpyxl.load_workbook(os.path.join(_fk, "out", "d_yok.xlsx"),
                                             data_only=True)[SAYFA]["Z97"].value
                      ).startswith("UYGUN DEĞİL"))
        for _ad, _g, _parca in _gecersiz:
            r.kontrol(f"[geçersiz girdi] {_ad} motorda reddediliyor",
                      not MK.hesapla(_g).get("aktif"))
            _gyol = os.path.join(_fk, "out", _ad + ".xlsx")
            if not os.path.isfile(_gyol):
                r.kontrol(f"[geçersiz girdi] {_ad} kitabı hesaplandı", False)
                continue
            _gws = openpyxl.load_workbook(_gyol, data_only=True)[SAYFA]
            for _z in ("Z242", "Z257", "Z271", "Z285"):
                _v = str(_gws[_z].value or "")
                r.kontrol(f"[geçersiz girdi] {_ad} · {_z} hüküm vermiyor",
                          _v.startswith("UYGUN DEĞİLDİR —") and _parca in _v,
                          f"→ {_v!r}")
        r.esit("[kanal f] her birleşim ayrı kitap olarak yeniden hesaplandı",
               len([a for a, _ in _fsen
                    if os.path.isfile(os.path.join(_fk, "out", a + ".xlsx"))]),
               len(_fsen))
        for _ad, _g in _fsen:
            _fyol = os.path.join(_fk, "out", _ad + ".xlsx")
            if not os.path.isfile(_fyol):
                continue
            _fws = openpyxl.load_workbook(_fyol, data_only=True)[SAYFA]
            _et = (f"{_g['kanal_sekli']} · {_g['kanal_isleme']}"
                   + ("  · ofis açıları değişik" if "_ofis" in _g else ""))
            _hm = MK.hesapla(_g)["_h"]
            for _h in _F_HUC:
                if _h in _hm:
                    r.kontrol(f"[kanal f] {_et} · {_h}", _esit(_hm[_h], _fws[_h].value),
                              f"→ motor {_hm[_h]!r}, kitap {_fws[_h].value!r}")
            #  ÖBÜR İŞLEMENİN SATIRI da doğru olmalı:  kitapta işleme sonradan
            #  değiştirilirse sınır o satırdan okunur.  Motor yalnız seçili
            #  işlemeninkini yazdığı için öbürü ayrı bir hesapla kıyaslanır.
            _obur = [x for x in MT.KANAL_ISLEME_SEKILLERI if x != _g["kanal_isleme"]][0]
            _gobur = MK.hesapla(dict(_g, kanal_isleme=_obur))
            if not _gobur.get("aktif"):
                #  Düz V'nin öbür işlemesi ( sertleştirilmemiş ) standart dışıdır
                #  ve reddedilir:  kıyaslanacak motor sonucu yoktur;  kitap o
                #  birleşimde zaten hüküm vermez ( bkz. [geçersiz girdi] ).
                continue
            _ho = _gobur["_h"]
            for _h in (("AJ198", "AL202") if _obur == "Sertleştirilmiş"
                       else ("AU206", "AV211")):
                r.kontrol(f"[kanal f] {_et} · öbür işleme {_h}",
                          _esit(_ho[_h], _fws[_h].value),
                          f"→ motor {_ho[_h]!r}, kitap {_fws[_h].value!r}")
            for _e in hata_hucresi_ara(_fyol):
                if _e.startswith(SAYFA):
                    r.kontrol(f"[kanal f] {_et} · Excel hata hücresi", False, _e)

    shutil.rmtree(GECICI, ignore_errors=True)
    return r


def _esit(a, b, tol=1e-7):
    try:
        return abs(float(a) - float(b)) <= tol * max(abs(float(a)), abs(float(b)), 1.0)
    except (TypeError, ValueError):
        return a == b


if __name__ == "__main__":
    sys.exit(0 if calistir().yazdir() else 1)
